from __future__ import annotations

import csv
import io
import json
import math
import os
import secrets
import shutil
import sqlite3
import threading
import uuid
from contextlib import contextmanager
from datetime import date, datetime, time, timedelta
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from functools import wraps
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo

from flask import Flask, Response, g, jsonify, make_response, render_template, request, send_file, session
from werkzeug.security import check_password_hash, generate_password_hash
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


ROOT = Path(__file__).resolve().parent
DATA_DIR = Path(os.environ.get("HOK_DATA_DIR", ROOT / "data")).resolve()
DB_PATH = DATA_DIR / "hok_filling_station.sqlite3"
BACKUP_DIR = DATA_DIR / "backups"
TIMEZONE_NAME = "Asia/Dhaka"
TZ = ZoneInfo(TIMEZONE_NAME)
BACKUP_LOCK = threading.Lock()

SCHEMA = r"""
PRAGMA foreign_keys = ON;
CREATE TABLE IF NOT EXISTS users (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    username TEXT NOT NULL UNIQUE COLLATE NOCASE,
    full_name TEXT NOT NULL,
    password_hash TEXT NOT NULL,
    role TEXT NOT NULL,
    is_active INTEGER NOT NULL DEFAULT 1,
    created_at TEXT NOT NULL,
    last_login TEXT
);
CREATE TABLE IF NOT EXISTS roles (name TEXT PRIMARY KEY, label TEXT NOT NULL);
CREATE TABLE IF NOT EXISTS role_permissions (
    role TEXT NOT NULL,
    permission TEXT NOT NULL,
    PRIMARY KEY (role, permission),
    FOREIGN KEY (role) REFERENCES roles(name) ON DELETE CASCADE
);
CREATE TABLE IF NOT EXISTS settings (key TEXT PRIMARY KEY, value TEXT NOT NULL);
CREATE TABLE IF NOT EXISTS products (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    name TEXT NOT NULL UNIQUE,
    category TEXT NOT NULL,
    unit TEXT NOT NULL,
    purchase_price_paisa INTEGER NOT NULL DEFAULT 0,
    selling_price_paisa INTEGER NOT NULL DEFAULT 0,
    current_stock REAL NOT NULL DEFAULT 0,
    min_stock REAL NOT NULL DEFAULT 0,
    active INTEGER NOT NULL DEFAULT 1,
    created_at TEXT NOT NULL,
    updated_at TEXT NOT NULL
);
CREATE TABLE IF NOT EXISTS price_history (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    product_id INTEGER NOT NULL,
    old_price_paisa INTEGER NOT NULL,
    new_price_paisa INTEGER NOT NULL,
    effective_at TEXT NOT NULL,
    user_id INTEGER,
    reason TEXT,
    FOREIGN KEY (product_id) REFERENCES products(id),
    FOREIGN KEY (user_id) REFERENCES users(id)
);
CREATE TABLE IF NOT EXISTS customers (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    name TEXT NOT NULL,
    phone TEXT,
    address TEXT,
    vehicle_no TEXT,
    customer_type TEXT NOT NULL DEFAULT 'Regular Customer',
    opening_due_paisa INTEGER NOT NULL DEFAULT 0,
    status TEXT NOT NULL DEFAULT 'active',
    created_at TEXT NOT NULL,
    updated_at TEXT NOT NULL
);
CREATE TABLE IF NOT EXISTS customer_transactions (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    customer_id INTEGER NOT NULL,
    date TEXT NOT NULL,
    type TEXT NOT NULL,
    amount_paisa INTEGER NOT NULL,
    source TEXT,
    source_id INTEGER,
    notes TEXT,
    user_id INTEGER,
    created_at TEXT NOT NULL,
    FOREIGN KEY (customer_id) REFERENCES customers(id),
    FOREIGN KEY (user_id) REFERENCES users(id)
);
CREATE TABLE IF NOT EXISTS suppliers (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    name TEXT NOT NULL,
    phone TEXT,
    address TEXT,
    company TEXT,
    opening_balance_paisa INTEGER NOT NULL DEFAULT 0,
    status TEXT NOT NULL DEFAULT 'active',
    created_at TEXT NOT NULL,
    updated_at TEXT NOT NULL
);
CREATE TABLE IF NOT EXISTS supplier_transactions (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    supplier_id INTEGER NOT NULL,
    date TEXT NOT NULL,
    type TEXT NOT NULL,
    amount_paisa INTEGER NOT NULL,
    source TEXT,
    source_id INTEGER,
    notes TEXT,
    user_id INTEGER,
    created_at TEXT NOT NULL,
    FOREIGN KEY (supplier_id) REFERENCES suppliers(id),
    FOREIGN KEY (user_id) REFERENCES users(id)
);
CREATE TABLE IF NOT EXISTS employees (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    name TEXT NOT NULL,
    phone TEXT,
    address TEXT,
    designation TEXT,
    join_date TEXT,
    salary_paisa INTEGER NOT NULL DEFAULT 0,
    status TEXT NOT NULL DEFAULT 'active',
    user_id INTEGER,
    created_at TEXT NOT NULL,
    updated_at TEXT NOT NULL,
    FOREIGN KEY (user_id) REFERENCES users(id)
);
CREATE TABLE IF NOT EXISTS attendance (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    employee_id INTEGER NOT NULL,
    date TEXT NOT NULL,
    shift_id INTEGER NOT NULL,
    status TEXT NOT NULL,
    check_in TEXT,
    check_out TEXT,
    notes TEXT,
    created_by INTEGER,
    created_at TEXT NOT NULL,
    updated_at TEXT NOT NULL,
    UNIQUE(employee_id,date,shift_id),
    FOREIGN KEY (employee_id) REFERENCES employees(id),
    FOREIGN KEY (shift_id) REFERENCES shifts(id),
    FOREIGN KEY (created_by) REFERENCES users(id)
);
CREATE TABLE IF NOT EXISTS pumps (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    pump_number TEXT NOT NULL UNIQUE,
    active INTEGER NOT NULL DEFAULT 1,
    created_at TEXT NOT NULL
);
CREATE TABLE IF NOT EXISTS nozzles (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    pump_id INTEGER NOT NULL,
    nozzle_number TEXT NOT NULL,
    product_id INTEGER NOT NULL,
    tank_id INTEGER,
    active INTEGER NOT NULL DEFAULT 1,
    last_meter REAL NOT NULL DEFAULT 0,
    UNIQUE (pump_id, nozzle_number),
    FOREIGN KEY (pump_id) REFERENCES pumps(id),
    FOREIGN KEY (product_id) REFERENCES products(id),
    FOREIGN KEY (tank_id) REFERENCES tanks(id)
);
CREATE TABLE IF NOT EXISTS shifts (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    name TEXT NOT NULL UNIQUE,
    start_time TEXT NOT NULL,
    end_time TEXT NOT NULL,
    active INTEGER NOT NULL DEFAULT 1,
    created_at TEXT NOT NULL
);
CREATE TABLE IF NOT EXISTS shift_closings (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    date TEXT NOT NULL,
    shift_id INTEGER NOT NULL,
    opening_cash_paisa INTEGER NOT NULL DEFAULT 0,
    closing_cash_paisa INTEGER NOT NULL DEFAULT 0,
    opening_meter REAL NOT NULL DEFAULT 0,
    closing_meter REAL NOT NULL DEFAULT 0,
    total_liters REAL NOT NULL DEFAULT 0,
    total_sales_paisa INTEGER NOT NULL DEFAULT 0,
    cash_sales_paisa INTEGER NOT NULL DEFAULT 0,
    credit_sales_paisa INTEGER NOT NULL DEFAULT 0,
    expenses_paisa INTEGER NOT NULL DEFAULT 0,
    deposits_paisa INTEGER NOT NULL DEFAULT 0,
    expected_cash_paisa INTEGER NOT NULL DEFAULT 0,
    cash_difference_paisa INTEGER NOT NULL DEFAULT 0,
    notes TEXT,
    closed_by INTEGER,
    closed_at TEXT NOT NULL,
    UNIQUE (date, shift_id),
    FOREIGN KEY (shift_id) REFERENCES shifts(id),
    FOREIGN KEY (closed_by) REFERENCES users(id)
);
CREATE TABLE IF NOT EXISTS tanks (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    tank_number TEXT NOT NULL UNIQUE,
    product_id INTEGER NOT NULL,
    capacity REAL NOT NULL DEFAULT 0,
    opening_quantity REAL NOT NULL DEFAULT 0,
    min_level REAL NOT NULL DEFAULT 0,
    max_level REAL NOT NULL DEFAULT 0,
    active INTEGER NOT NULL DEFAULT 1,
    created_at TEXT NOT NULL,
    FOREIGN KEY (product_id) REFERENCES products(id)
);
CREATE TABLE IF NOT EXISTS tank_movements (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    tank_id INTEGER NOT NULL,
    date TEXT NOT NULL,
    kind TEXT NOT NULL,
    quantity REAL NOT NULL,
    reference_type TEXT,
    reference_id INTEGER,
    notes TEXT,
    user_id INTEGER,
    created_at TEXT NOT NULL,
    FOREIGN KEY (tank_id) REFERENCES tanks(id),
    FOREIGN KEY (user_id) REFERENCES users(id)
);
CREATE TABLE IF NOT EXISTS sales (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    memo_no TEXT NOT NULL UNIQUE,
    date TEXT NOT NULL,
    time TEXT NOT NULL,
    customer_id INTEGER,
    vehicle_no TEXT,
    pump_id INTEGER,
    nozzle_id INTEGER,
    shift_id INTEGER,
    payment_type TEXT NOT NULL,
    total_paisa INTEGER NOT NULL,
    paid_cash_paisa INTEGER NOT NULL DEFAULT 0,
    paid_bank_paisa INTEGER NOT NULL DEFAULT 0,
    paid_mobile_paisa INTEGER NOT NULL DEFAULT 0,
    credit_paisa INTEGER NOT NULL DEFAULT 0,
    bank_account_id INTEGER,
    operator_id INTEGER,
    user_id INTEGER,
    notes TEXT,
    status TEXT NOT NULL DEFAULT 'posted',
    revision INTEGER NOT NULL DEFAULT 1,
    created_at TEXT NOT NULL,
    updated_at TEXT NOT NULL,
    FOREIGN KEY (customer_id) REFERENCES customers(id),
    FOREIGN KEY (pump_id) REFERENCES pumps(id),
    FOREIGN KEY (nozzle_id) REFERENCES nozzles(id),
    FOREIGN KEY (shift_id) REFERENCES shifts(id),
    FOREIGN KEY (bank_account_id) REFERENCES bank_accounts(id),
    FOREIGN KEY (operator_id) REFERENCES employees(id),
    FOREIGN KEY (user_id) REFERENCES users(id)
);
CREATE TABLE IF NOT EXISTS sale_items (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    sale_id INTEGER NOT NULL,
    product_id INTEGER NOT NULL,
    quantity REAL NOT NULL,
    unit_price_paisa INTEGER NOT NULL,
    line_total_paisa INTEGER NOT NULL,
    unit_cost_paisa INTEGER NOT NULL DEFAULT 0,
    FOREIGN KEY (sale_id) REFERENCES sales(id),
    FOREIGN KEY (product_id) REFERENCES products(id)
);
CREATE TABLE IF NOT EXISTS sale_revisions (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    sale_id INTEGER NOT NULL,
    revision INTEGER NOT NULL,
    changed_at TEXT NOT NULL,
    user_id INTEGER,
    snapshot TEXT NOT NULL,
    reason TEXT,
    FOREIGN KEY (sale_id) REFERENCES sales(id),
    FOREIGN KEY (user_id) REFERENCES users(id)
);
CREATE TABLE IF NOT EXISTS purchases (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    invoice_no TEXT NOT NULL UNIQUE,
    date TEXT NOT NULL,
    supplier_id INTEGER NOT NULL,
    product_id INTEGER NOT NULL,
    quantity REAL NOT NULL,
    purchase_rate_paisa INTEGER NOT NULL,
    total_paisa INTEGER NOT NULL,
    paid_amount_paisa INTEGER NOT NULL DEFAULT 0,
    due_amount_paisa INTEGER NOT NULL DEFAULT 0,
    transport_cost_paisa INTEGER NOT NULL DEFAULT 0,
    other_cost_paisa INTEGER NOT NULL DEFAULT 0,
    payment_method TEXT NOT NULL DEFAULT 'cash',
    bank_account_id INTEGER,
    tank_id INTEGER,
    challan_no TEXT,
    notes TEXT,
    status TEXT NOT NULL DEFAULT 'posted',
    user_id INTEGER,
    created_at TEXT NOT NULL,
    FOREIGN KEY (supplier_id) REFERENCES suppliers(id),
    FOREIGN KEY (product_id) REFERENCES products(id),
    FOREIGN KEY (tank_id) REFERENCES tanks(id),
    FOREIGN KEY (bank_account_id) REFERENCES bank_accounts(id),
    FOREIGN KEY (user_id) REFERENCES users(id)
);
CREATE TABLE IF NOT EXISTS stock_movements (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    product_id INTEGER NOT NULL,
    kind TEXT NOT NULL,
    quantity REAL NOT NULL,
    business_date TEXT NOT NULL,
    unit_cost_paisa INTEGER NOT NULL DEFAULT 0,
    reference_type TEXT,
    reference_id INTEGER,
    notes TEXT,
    user_id INTEGER,
    created_at TEXT NOT NULL,
    FOREIGN KEY (product_id) REFERENCES products(id),
    FOREIGN KEY (user_id) REFERENCES users(id)
);
CREATE TABLE IF NOT EXISTS meter_readings (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    date TEXT NOT NULL,
    shift_id INTEGER,
    pump_id INTEGER NOT NULL,
    nozzle_id INTEGER NOT NULL,
    product_id INTEGER NOT NULL,
    opening_meter REAL NOT NULL,
    closing_meter REAL NOT NULL,
    sales_liters REAL NOT NULL,
    rate_paisa INTEGER NOT NULL,
    total_paisa INTEGER NOT NULL,
    operator_id INTEGER,
    user_id INTEGER,
    notes TEXT,
    created_at TEXT NOT NULL,
    FOREIGN KEY (shift_id) REFERENCES shifts(id),
    FOREIGN KEY (pump_id) REFERENCES pumps(id),
    FOREIGN KEY (nozzle_id) REFERENCES nozzles(id),
    FOREIGN KEY (product_id) REFERENCES products(id),
    FOREIGN KEY (operator_id) REFERENCES employees(id),
    FOREIGN KEY (user_id) REFERENCES users(id)
);
CREATE TABLE IF NOT EXISTS tank_dips (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    date TEXT NOT NULL,
    tank_id INTEGER NOT NULL,
    dip_reading REAL NOT NULL,
    estimated_quantity REAL NOT NULL,
    book_stock REAL NOT NULL,
    physical_stock REAL NOT NULL,
    difference REAL NOT NULL,
    operator_id INTEGER,
    verified INTEGER NOT NULL DEFAULT 0,
    verified_by INTEGER,
    notes TEXT,
    user_id INTEGER,
    created_at TEXT NOT NULL,
    FOREIGN KEY (tank_id) REFERENCES tanks(id),
    FOREIGN KEY (operator_id) REFERENCES employees(id),
    FOREIGN KEY (verified_by) REFERENCES users(id),
    FOREIGN KEY (user_id) REFERENCES users(id)
);
CREATE TABLE IF NOT EXISTS expenses (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    date TEXT NOT NULL,
    category TEXT NOT NULL,
    description TEXT,
    amount_paisa INTEGER NOT NULL,
    payment_method TEXT NOT NULL,
    bank_account_id INTEGER,
    paid_to TEXT,
    voucher_no TEXT,
    shift_id INTEGER,
    status TEXT NOT NULL DEFAULT 'approved',
    created_by INTEGER,
    approved_by INTEGER,
    created_at TEXT NOT NULL,
    FOREIGN KEY (bank_account_id) REFERENCES bank_accounts(id),
    FOREIGN KEY (created_by) REFERENCES users(id),
    FOREIGN KEY (approved_by) REFERENCES users(id)
);
CREATE TABLE IF NOT EXISTS bank_accounts (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    bank_name TEXT NOT NULL,
    account_name TEXT NOT NULL,
    account_number TEXT,
    opening_balance_paisa INTEGER NOT NULL DEFAULT 0,
    status TEXT NOT NULL DEFAULT 'active',
    created_at TEXT NOT NULL,
    updated_at TEXT NOT NULL
);
CREATE TABLE IF NOT EXISTS bank_transactions (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    date TEXT NOT NULL,
    bank_id INTEGER,
    channel TEXT NOT NULL DEFAULT 'bank',
    direction TEXT NOT NULL,
    type TEXT NOT NULL,
    amount_paisa INTEGER NOT NULL,
    slip_no TEXT,
    depositor TEXT,
    reference TEXT,
    notes TEXT,
    reference_type TEXT,
    reference_id INTEGER,
    shift_id INTEGER,
    user_id INTEGER,
    created_at TEXT NOT NULL,
    FOREIGN KEY (shift_id) REFERENCES shifts(id),
    FOREIGN KEY (bank_id) REFERENCES bank_accounts(id),
    FOREIGN KEY (user_id) REFERENCES users(id)
);
CREATE TABLE IF NOT EXISTS cash_transactions (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    date TEXT NOT NULL,
    direction TEXT NOT NULL,
    type TEXT NOT NULL,
    amount_paisa INTEGER NOT NULL,
    reference_type TEXT,
    reference_id INTEGER,
    notes TEXT,
    shift_id INTEGER,
    user_id INTEGER,
    created_at TEXT NOT NULL,
    FOREIGN KEY (shift_id) REFERENCES shifts(id),
    FOREIGN KEY (user_id) REFERENCES users(id)
);
CREATE TABLE IF NOT EXISTS audit_logs (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    user_id INTEGER,
    timestamp TEXT NOT NULL,
    action TEXT NOT NULL,
    module TEXT NOT NULL,
    record_id TEXT,
    previous_value TEXT,
    new_value TEXT,
    reason TEXT,
    FOREIGN KEY (user_id) REFERENCES users(id)
);
CREATE TABLE IF NOT EXISTS backup_history (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    filename TEXT NOT NULL,
    path TEXT NOT NULL,
    status TEXT NOT NULL,
    backup_type TEXT NOT NULL DEFAULT 'manual',
    size_bytes INTEGER NOT NULL DEFAULT 0,
    created_by INTEGER,
    created_at TEXT NOT NULL,
    message TEXT,
    FOREIGN KEY (created_by) REFERENCES users(id)
);
CREATE INDEX IF NOT EXISTS idx_sales_date_status ON sales(date, status);
CREATE INDEX IF NOT EXISTS idx_sales_customer ON sales(customer_id, date);
CREATE INDEX IF NOT EXISTS idx_sale_items_product ON sale_items(product_id, sale_id);
CREATE INDEX IF NOT EXISTS idx_stock_product_time ON stock_movements(product_id, created_at);
CREATE INDEX IF NOT EXISTS idx_customer_txn ON customer_transactions(customer_id, date);
CREATE INDEX IF NOT EXISTS idx_supplier_txn ON supplier_transactions(supplier_id, date);
CREATE INDEX IF NOT EXISTS idx_meter_nozzle_date ON meter_readings(nozzle_id, date, id);
CREATE INDEX IF NOT EXISTS idx_expenses_date ON expenses(date, status);
CREATE INDEX IF NOT EXISTS idx_cash_date ON cash_transactions(date);
CREATE INDEX IF NOT EXISTS idx_bank_date ON bank_transactions(date);
CREATE INDEX IF NOT EXISTS idx_audit_time ON audit_logs(timestamp);
"""

ROLE_LABELS = {
    "owner": "Owner / Admin",
    "manager": "Manager",
    "accountant": "Accountant",
    "computer_operator": "Computer Operator",
    "sales_operator": "Sales Operator",
}

PERMISSION_LABELS = {
    "dashboard:view": "View dashboard",
    "sales:view": "View sales and memos",
    "sales:create": "Create sales",
    "sales:edit": "Edit posted sales",
    "sales:cancel": "Cancel / reverse sales",
    "products:view": "View products and prices",
    "products:manage": "Manage products",
    "prices:manage": "Change selling prices",
    "inventory:view": "View stock and inventory",
    "inventory:adjust": "Adjust stock",
    "purchases:view": "View purchases",
    "purchases:create": "Create purchases",
    "suppliers:manage": "Manage suppliers and payments",
    "customers:manage": "Manage customers",
    "due:collect": "Collect customer due",
    "cashbook:view": "View cashbook",
    "cashbook:manage": "Record cash income / withdrawal",
    "banks:manage": "Manage bank accounts and deposits",
    "expenses:manage": "Manage and approve expenses",
    "employees:manage": "Manage employees",
    "shifts:manage": "Manage and close shifts",
    "pumps:manage": "Manage pumps and nozzles",
    "meters:manage": "Record meter readings",
    "tanks:manage": "Manage tanks and dip readings",
    "reports:view": "View and export reports",
    "notifications:view": "View notifications",
    "audit:view": "View audit logs",
    "users:manage": "Manage users and roles",
    "settings:manage": "Manage business settings",
    "backup:manage": "Create / restore backups",
}

DEFAULT_ROLE_PERMISSIONS = {
    "owner": list(PERMISSION_LABELS.keys()),
    "manager": [
        "dashboard:view", "sales:view", "sales:create", "sales:edit", "sales:cancel",
        "products:view", "inventory:view", "inventory:adjust", "purchases:view", "purchases:create",
        "suppliers:manage", "customers:manage", "due:collect", "cashbook:view", "cashbook:manage",
        "banks:manage", "expenses:manage", "employees:manage", "shifts:manage", "pumps:manage",
        "meters:manage", "tanks:manage", "reports:view", "notifications:view", "prices:manage",
    ],
    "accountant": [
        "dashboard:view", "sales:view", "products:view", "inventory:view", "purchases:view",
        "suppliers:manage", "customers:manage", "due:collect", "cashbook:view", "cashbook:manage",
        "banks:manage", "expenses:manage", "reports:view", "notifications:view",
    ],
    "computer_operator": [
        "dashboard:view", "sales:view", "sales:create", "products:view", "inventory:view",
        "customers:manage", "due:collect", "meters:manage", "pumps:manage", "shifts:manage",
        "reports:view", "notifications:view",
    ],
    "sales_operator": [
        "dashboard:view", "sales:view", "sales:create", "products:view", "customers:manage",
        "due:collect", "meters:manage", "shifts:manage", "notifications:view",
    ],
}

DEFAULT_SETTINGS = {
    "business_name": "মেসার্স হক ফিলিং স্টেশন",
    "dealer_name": "যমুনা অয়েল কোম্পানি লিঃ",
    "business_address": "খোয়ারপাড়, শেরপুর টাউন, শেরপুর-২১০০",
    "country": "Bangladesh",
    "currency": "৳",
    "timezone": "Asia/Dhaka",
    "language": "বাংলা",
    "contact_phone": "",
    "contact_mobile": "",
    "logo_data": "",
    "memo_prefix": "",
    "memo_start": "80681",
    "memo_next": "80681",
    "memo_format": "{prefix}{number}",
    "purchase_prefix": "PUR-",
    "purchase_next": "1",
    "opening_cash": "0.00",
    "due_warning": "50000.00",
    "backup_auto": "1",
    "last_auto_backup": "",
}

INITIAL_PRODUCTS = [
    ("ডিজেল", "fuel", "লিটার", 0, 13600, 0),
    ("পেট্রোল", "fuel", "লিটার", 0, 16100, 0),
    ("অকটেন", "fuel", "লিটার", 0, 16600, 0),
    ("মবিল", "lubricant", "লিটার", 0, 30000, 0),
    ("মবিল স্পেশাল", "lubricant", "বোতল", 0, 71500, 0),
    ("মবিল সুপার 4T", "lubricant", "বোতল", 0, 0, 0),
    ("গিয়ার অয়েল", "lubricant", "লিটার", 0, 0, 0),
    ("ব্রেক অয়েল", "lubricant", "বোতল", 0, 0, 0),
    ("গ্রীজ", "lubricant", "কেজি", 0, 0, 0),
    ("ভিসকো কুলিং ওয়াটার", "lubricant", "বোতল", 0, 0, 0),
    ("ফ্লাশিং অয়েল", "lubricant", "লিটার", 0, 0, 0),
    ("অন্যান্য", "other", "পিস", 0, 0, 0),
]

EXPENSE_CATEGORIES = [
    "Salary", "Electricity", "Generator", "Maintenance", "Transport", "Office Expense",
    "Internet", "Telephone", "Rent", "Bank Charge", "Cleaning", "Stationery", "Repair", "Other",
]
CUSTOMER_TYPES = ["Regular Customer", "Business Customer", "Transport Customer", "Credit Customer", "Other"]
PAYMENT_METHODS = ["cash", "credit", "bank", "mobile", "mixed"]


class ApiError(Exception):
    def __init__(self, message: str, status: int = 400, code: str = "invalid_request"):
        super().__init__(message)
        self.message, self.status, self.code = message, status, code


def now_dt() -> datetime:
    return datetime.now(TZ)


def now_text() -> str:
    return now_dt().strftime("%Y-%m-%d %H:%M:%S")


def today_text() -> str:
    return now_dt().strftime("%Y-%m-%d")


def money_minor(value: Any, field: str = "amount", allow_negative: bool = False) -> int:
    try:
        amount = Decimal(str(value if value not in (None, "") else 0)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    except (InvalidOperation, ValueError, TypeError):
        raise ApiError(f"{field} must be a valid amount.")
    if not amount.is_finite() or (not allow_negative and amount < 0):
        raise ApiError(f"{field} must be a non-negative amount.")
    return int(amount * 100)


def to_money(value: Any) -> float:
    return round(int(value or 0) / 100, 2)


def quantity(value: Any, field: str = "quantity", allow_negative: bool = False) -> float:
    try:
        amount = Decimal(str(value if value not in (None, "") else 0)).quantize(Decimal("0.001"), rounding=ROUND_HALF_UP)
    except (InvalidOperation, ValueError, TypeError):
        raise ApiError(f"{field} must be a valid number.")
    if not amount.is_finite() or (not allow_negative and amount < 0):
        raise ApiError(f"{field} must be a non-negative number.")
    return float(amount)


def meter_value(value: Any, field: str) -> float:
    return quantity(value, field)


def amount_for_quantity(qty: float, unit_price_paisa: int) -> int:
    return int((Decimal(str(qty)) * Decimal(int(unit_price_paisa))).quantize(Decimal("1"), rounding=ROUND_HALF_UP))


def parse_date(value: Any, field: str = "date", default_today: bool = True) -> str:
    if not value:
        if default_today:
            return today_text()
        raise ApiError(f"{field} is required.")
    text = str(value).strip()
    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(text, fmt).strftime("%Y-%m-%d")
        except ValueError:
            pass
    raise ApiError(f"{field} must be in YYYY-MM-DD or DD-MM-YYYY format.")


def parse_int(value: Any, field: str, default: int | None = None) -> int | None:
    if value in (None, ""):
        return default
    try:
        number = int(value)
    except (TypeError, ValueError):
        raise ApiError(f"{field} must be a whole number.")
    return number


def db_connection(path: Path = DB_PATH) -> sqlite3.Connection:
    conn = sqlite3.connect(str(path), timeout=30, isolation_level="DEFERRED")
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys = ON")
    conn.execute("PRAGMA busy_timeout = 30000")
    return conn


def get_db() -> sqlite3.Connection:
    if "db" not in g:
        g.db = db_connection()
    return g.db


@contextmanager
def atomic():
    conn = get_db()
    conn.execute("BEGIN IMMEDIATE")
    try:
        yield conn
        conn.commit()
    except Exception:
        conn.rollback()
        raise


def fetchone(sql: str, params: tuple = (), conn: sqlite3.Connection | None = None):
    return (conn or get_db()).execute(sql, params).fetchone()


def fetchall(sql: str, params: tuple = (), conn: sqlite3.Connection | None = None):
    return (conn or get_db()).execute(sql, params).fetchall()


def rowdict(row: sqlite3.Row | None) -> dict | None:
    return dict(row) if row is not None else None


def request_data() -> dict:
    data = request.get_json(silent=True)
    if not isinstance(data, dict):
        raise ApiError("A JSON request body is required.")
    return data


def setting(key: str, default: str = "") -> str:
    row = fetchone("SELECT value FROM settings WHERE key = ?", (key,))
    return row["value"] if row else default


def set_setting(conn: sqlite3.Connection, key: str, value: Any) -> None:
    conn.execute("INSERT INTO settings(key, value) VALUES(?, ?) ON CONFLICT(key) DO UPDATE SET value=excluded.value", (key, str(value)))


def settings_dict() -> dict:
    values = dict(DEFAULT_SETTINGS)
    values.update({row["key"]: row["value"] for row in fetchall("SELECT key, value FROM settings")})
    return values


def audit(conn: sqlite3.Connection, action: str, module: str, record_id: Any = None,
          previous: Any = None, new: Any = None, reason: str | None = None) -> None:
    user_id = session.get("uid")
    conn.execute(
        "INSERT INTO audit_logs(user_id, timestamp, action, module, record_id, previous_value, new_value, reason) VALUES(?,?,?,?,?,?,?,?)",
        (user_id, now_text(), action, module, None if record_id is None else str(record_id),
         json.dumps(previous, ensure_ascii=False, default=str) if previous is not None else None,
         json.dumps(new, ensure_ascii=False, default=str) if new is not None else None, reason),
    )


def ensure_schema_migrations(conn: sqlite3.Connection) -> None:
    """Small, idempotent forward migrations for installations already using this app."""
    additions = {
        "nozzles": [("tank_id", "INTEGER REFERENCES tanks(id)")],
        "sales": [("revision", "INTEGER NOT NULL DEFAULT 1")],
        "stock_movements": [("business_date", "TEXT")],
        "cash_transactions": [("shift_id", "INTEGER REFERENCES shifts(id)")],
        "bank_transactions": [("shift_id", "INTEGER REFERENCES shifts(id)")],
        "expenses": [("shift_id", "INTEGER REFERENCES shifts(id)")],
        "shift_closings": [("opening_meter", "REAL NOT NULL DEFAULT 0"), ("closing_meter", "REAL NOT NULL DEFAULT 0"), ("total_liters", "REAL NOT NULL DEFAULT 0"), ("total_sales_paisa", "INTEGER NOT NULL DEFAULT 0")],
    }
    for table, columns in additions.items():
        known = {row["name"] for row in conn.execute(f"PRAGMA table_info({table})")}
        for name, definition in columns:
            if name not in known:
                conn.execute(f"ALTER TABLE {table} ADD COLUMN {name} {definition}")
    if "business_date" in {row["name"] for row in conn.execute("PRAGMA table_info(stock_movements)")}:
        conn.execute("UPDATE stock_movements SET business_date=substr(created_at,1,10) WHERE business_date IS NULL OR business_date=''")


def init_database() -> None:
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    BACKUP_DIR.mkdir(parents=True, exist_ok=True)
    conn = db_connection()
    try:
        conn.executescript(SCHEMA)
        ensure_schema_migrations(conn)
        for key, value in DEFAULT_SETTINGS.items():
            conn.execute("INSERT OR IGNORE INTO settings(key, value) VALUES(?, ?)", (key, value))
        for role, label in ROLE_LABELS.items():
            conn.execute("INSERT OR IGNORE INTO roles(name, label) VALUES(?, ?)", (role, label))
            for permission in DEFAULT_ROLE_PERMISSIONS[role]:
                conn.execute("INSERT OR IGNORE INTO role_permissions(role, permission) VALUES(?, ?)", (role, permission))
        stamp = now_text()
        for name, category, unit, purchase, selling, min_stock in INITIAL_PRODUCTS:
            conn.execute(
                "INSERT OR IGNORE INTO products(name, category, unit, purchase_price_paisa, selling_price_paisa, current_stock, min_stock, active, created_at, updated_at) VALUES(?,?,?,?,?,0,?,1,?,?)",
                (name, category, unit, purchase, selling, min_stock, stamp, stamp),
            )
        conn.execute("INSERT OR IGNORE INTO shifts(name,start_time,end_time,active,created_at) VALUES('Day Shift','09:00','21:00',1,?)", (stamp,))
        conn.execute("INSERT OR IGNORE INTO shifts(name,start_time,end_time,active,created_at) VALUES('Night Shift','21:00','09:00',1,?)", (stamp,))
        conn.commit()
        conn.execute("PRAGMA journal_mode = WAL")
    finally:
        conn.close()


def secure_secret() -> str:
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    env_secret = os.environ.get("HOK_SECRET_KEY")
    if env_secret:
        return env_secret
    key_file = DATA_DIR / "session.key"
    if key_file.exists():
        return key_file.read_text(encoding="utf-8").strip()
    secret = secrets.token_urlsafe(48)
    key_file.write_text(secret, encoding="utf-8")
    try:
        key_file.chmod(0o600)
    except OSError:
        pass
    return secret


def current_user() -> sqlite3.Row | None:
    uid = session.get("uid")
    if not uid:
        return None
    return fetchone("SELECT id, username, full_name, role, is_active FROM users WHERE id=?", (uid,))


def user_permissions(user: sqlite3.Row | None = None) -> list[str]:
    user = user or current_user()
    if not user:
        return []
    return [row["permission"] for row in fetchall("SELECT permission FROM role_permissions WHERE role=? ORDER BY permission", (user["role"],))]


def permission_required(permission: str):
    def decorator(func):
        @wraps(func)
        def wrapped(*args, **kwargs):
            user = current_user()
            if not user or not user["is_active"]:
                raise ApiError("Please sign in to continue.", 401, "authentication_required")
            if permission not in user_permissions(user):
                raise ApiError("You do not have permission to perform this action.", 403, "permission_denied")
            g.user = user
            return func(*args, **kwargs)
        return wrapped
    return decorator


def auth_required(func):
    @wraps(func)
    def wrapped(*args, **kwargs):
        user = current_user()
        if not user or not user["is_active"]:
            raise ApiError("Please sign in to continue.", 401, "authentication_required")
        g.user = user
        return func(*args, **kwargs)
    return wrapped


def active_record(table: str, record_id: int, message: str = "Record not found."):
    allowed = {"products", "customers", "suppliers", "employees", "pumps", "nozzles", "shifts", "tanks", "bank_accounts"}
    if table not in allowed:
        raise RuntimeError("Unsafe table")
    row = fetchone(f"SELECT * FROM {table} WHERE id=?", (record_id,))
    if not row:
        raise ApiError(message, 404, "not_found")
    return row


def ensure_positive(value: int | float, field: str) -> None:
    if value <= 0:
        raise ApiError(f"{field} must be greater than zero.")


def date_range_params() -> tuple[str, str]:
    start = parse_date(request.args.get("start"), "start date") if request.args.get("start") else today_text()
    end = parse_date(request.args.get("end"), "end date") if request.args.get("end") else today_text()
    if start > end:
        raise ApiError("Start date must be on or before end date.")
    return start, end


def format_memo_number(value: int) -> str:
    prefix = setting("memo_prefix")
    fmt = setting("memo_format", "{prefix}{number}")
    try:
        return fmt.format(prefix=prefix, number=value)
    except (KeyError, ValueError, IndexError):
        return f"{prefix}{value}"


def next_number(conn: sqlite3.Connection, key: str) -> int:
    try:
        number = max(1, int(setting(key, "1")))
    except ValueError:
        number = 1
    set_setting(conn, key, number + 1)
    return number


def audit_dict(row: sqlite3.Row) -> dict:
    result = dict(row)
    for key in ("previous_value", "new_value"):
        if result.get(key):
            try:
                result[key] = json.loads(result[key])
            except (TypeError, ValueError):
                pass
    return result


def balance_customer(customer_id: int, conn: sqlite3.Connection | None = None) -> int:
    row = fetchone("SELECT opening_due_paisa FROM customers WHERE id=?", (customer_id,), conn)
    if not row:
        return 0
    due = fetchone("SELECT COALESCE(SUM(amount_paisa),0) AS amount FROM customer_transactions WHERE customer_id=?", (customer_id,), conn)
    return int(row["opening_due_paisa"] or 0) + int(due["amount"] or 0)


def balance_supplier(supplier_id: int, conn: sqlite3.Connection | None = None) -> int:
    row = fetchone("SELECT opening_balance_paisa FROM suppliers WHERE id=?", (supplier_id,), conn)
    if not row:
        return 0
    due = fetchone("SELECT COALESCE(SUM(amount_paisa),0) AS amount FROM supplier_transactions WHERE supplier_id=?", (supplier_id,), conn)
    return int(row["opening_balance_paisa"] or 0) + int(due["amount"] or 0)


def cash_balance(conn: sqlite3.Connection | None = None) -> int:
    rows = fetchall("SELECT direction, COALESCE(SUM(amount_paisa),0) AS amount FROM cash_transactions GROUP BY direction", (), conn)
    net = sum((1 if row["direction"] == "in" else -1) * int(row["amount"] or 0) for row in rows)
    return money_minor(setting("opening_cash", "0")) + net


def tank_book_quantity(tank_id: int, conn: sqlite3.Connection | None = None) -> float:
    row = fetchone("SELECT opening_quantity FROM tanks WHERE id=?", (tank_id,), conn)
    if not row:
        return 0.0
    moves = fetchone("SELECT COALESCE(SUM(quantity),0) value FROM tank_movements WHERE tank_id=?", (tank_id,), conn)
    return round(float(row["opening_quantity"] or 0) + float(moves["value"] or 0), 3)


def bank_balance(bank_id: int, conn: sqlite3.Connection | None = None) -> int:
    row = fetchone("SELECT opening_balance_paisa FROM bank_accounts WHERE id=?", (bank_id,), conn)
    if not row:
        return 0
    tx = fetchall("SELECT direction, COALESCE(SUM(amount_paisa),0) AS amount FROM bank_transactions WHERE bank_id=? GROUP BY direction", (bank_id,), conn)
    net = sum((1 if item["direction"] == "in" else -1) * int(item["amount"] or 0) for item in tx)
    return int(row["opening_balance_paisa"] or 0) + net


app = Flask(__name__, static_folder="static", template_folder="templates")
app.config.update(
    SECRET_KEY=secure_secret(),
    SESSION_COOKIE_HTTPONLY=True,
    SESSION_COOKIE_SAMESITE="Strict",
    SESSION_COOKIE_SECURE=os.environ.get("HOK_COOKIE_SECURE", "0").lower() in ("1", "true", "yes"),
    PERMANENT_SESSION_LIFETIME=timedelta(hours=10),
    MAX_CONTENT_LENGTH=2 * 1024 * 1024,
)

init_database()


@app.teardown_appcontext
def close_db(_error=None):
    conn = g.pop("db", None)
    if conn is not None:
        conn.close()


@app.before_request
def enforce_csrf_and_auto_backup():
    if request.path.startswith("/api/") and request.method in {"POST", "PUT", "PATCH", "DELETE"}:
        exempt = {"/api/setup", "/api/auth/login"}
        if request.path not in exempt and session.get("uid"):
            supplied = request.headers.get("X-CSRF-Token", "")
            expected = session.get("csrf_token", "")
            if not supplied or not expected or not secrets.compare_digest(supplied, expected):
                raise ApiError("Your session token expired. Refresh the page and try again.", 400, "csrf_failed")
    if request.path.startswith("/api/") and request.path not in {"/api/setup/status", "/api/auth/me", "/api/auth/login", "/api/setup"}:
        maybe_automatic_backup()


def maybe_automatic_backup() -> None:
    if not BACKUP_LOCK.acquire(blocking=False):
        return
    try:
        conn = db_connection()
        try:
            if not conn.execute("SELECT 1 FROM users LIMIT 1").fetchone():
                return
            values = {row["key"]: row["value"] for row in conn.execute("SELECT key,value FROM settings WHERE key IN ('backup_auto','last_auto_backup')")}
            if values.get("backup_auto", "1") != "1":
                return
            last = values.get("last_auto_backup", "")
            if last:
                try:
                    if now_dt() - datetime.strptime(last, "%Y-%m-%d %H:%M:%S").replace(tzinfo=TZ) < timedelta(hours=24):
                        return
                except ValueError:
                    pass
            filename = f"hok-auto-{now_dt().strftime('%Y%m%d-%H%M%S')}.sqlite3"
            path = BACKUP_DIR / filename
            dest = sqlite3.connect(str(path))
            try:
                conn.backup(dest)
            finally:
                dest.close()
            size = path.stat().st_size
            stamp = now_text()
            conn.execute("INSERT INTO backup_history(filename,path,status,backup_type,size_bytes,created_at,message) VALUES(?,?,'success','automatic',?,?,?)",
                         (filename, str(path), size, stamp, "Automatic daily backup"))
            conn.execute("UPDATE settings SET value=? WHERE key='last_auto_backup'", (stamp,))
            conn.commit()
        except Exception as exc:
            try:
                conn.execute("INSERT INTO backup_history(filename,path,status,backup_type,size_bytes,created_at,message) VALUES(?,?,'failed','automatic',0,?,?)",
                             ("automatic-backup", str(BACKUP_DIR), now_text(), str(exc)[:500]))
                conn.commit()
            except Exception:
                pass
        finally:
            conn.close()
    finally:
        BACKUP_LOCK.release()


@app.errorhandler(ApiError)
def handle_api_error(error: ApiError):
    return jsonify({"error": error.message, "code": error.code}), error.status


@app.errorhandler(404)
def handle_not_found(_error):
    if request.path.startswith("/api/"):
        return jsonify({"error": "The requested resource was not found.", "code": "not_found"}), 404
    return "Not found", 404


@app.get("/")
def index():
    return render_template("index.html")


@app.get("/api/setup/status")
def setup_status():
    row = fetchone("SELECT COUNT(*) AS count FROM users")
    return jsonify({"setup_required": int(row["count"]) == 0, "business_name": setting("business_name")})


@app.post("/api/setup")
def setup_owner():
    if fetchone("SELECT 1 FROM users LIMIT 1"):
        raise ApiError("Initial setup has already been completed.", 409, "setup_complete")
    data = request_data()
    username = str(data.get("username", "")).strip()
    full_name = str(data.get("full_name", "")).strip()
    password = str(data.get("password", ""))
    if len(username) < 3 or len(username) > 40:
        raise ApiError("Username must be between 3 and 40 characters.")
    if not full_name:
        raise ApiError("Please enter the owner's full name.")
    if len(password) < 10:
        raise ApiError("Choose a password with at least 10 characters.")
    with atomic() as conn:
        cur = conn.execute("INSERT INTO users(username,full_name,password_hash,role,is_active,created_at) VALUES(?,?,?,?,1,?)",
                           (username, full_name, generate_password_hash(password, method="scrypt"), "owner", now_text()))
        session.clear()
        session["uid"] = cur.lastrowid
        session["csrf_token"] = secrets.token_urlsafe(32)
        session.permanent = True
        audit(conn, "initial_setup", "users", cur.lastrowid, None, {"username": username, "role": "owner"})
    return jsonify({"ok": True, "csrf_token": session["csrf_token"], "user": {"id": cur.lastrowid, "username": username, "full_name": full_name, "role": "owner"}})


@app.post("/api/auth/login")
def login():
    data = request_data()
    username = str(data.get("username", "")).strip()
    password = str(data.get("password", ""))
    user = fetchone("SELECT * FROM users WHERE username=? COLLATE NOCASE", (username,))
    if not user or not user["is_active"] or not check_password_hash(user["password_hash"], password):
        raise ApiError("Incorrect username or password.", 401, "invalid_credentials")
    session.clear()
    session["uid"] = user["id"]
    session["csrf_token"] = secrets.token_urlsafe(32)
    session.permanent = True
    with atomic() as conn:
        conn.execute("UPDATE users SET last_login=? WHERE id=?", (now_text(), user["id"]))
        conn.execute("INSERT INTO audit_logs(user_id,timestamp,action,module,record_id,new_value) VALUES(?,?,?,?,?,?)",
                     (user["id"], now_text(), "login", "auth", str(user["id"]), json.dumps({"username": username})))
    return jsonify({"ok": True, "csrf_token": session["csrf_token"], "user": {"id": user["id"], "username": user["username"], "full_name": user["full_name"], "role": user["role"]}})


@app.post("/api/auth/logout")
@auth_required
def logout():
    user_id = g.user["id"]
    with atomic() as conn:
        audit(conn, "logout", "auth", user_id)
    session.clear()
    return jsonify({"ok": True})


@app.get("/api/auth/me")
def auth_me():
    user = current_user()
    if not user or not user["is_active"]:
        return jsonify({"authenticated": False, "csrf_token": ""})
    if not session.get("csrf_token"):
        session["csrf_token"] = secrets.token_urlsafe(32)
    return jsonify({"authenticated": True, "csrf_token": session["csrf_token"], "user": {"id": user["id"], "username": user["username"], "full_name": user["full_name"], "role": user["role"]}, "permissions": user_permissions(user)})


@app.get("/api/bootstrap")
@auth_required
def bootstrap():
    user = g.user
    return jsonify({
        "user": {"id": user["id"], "username": user["username"], "full_name": user["full_name"], "role": user["role"]},
        "permissions": user_permissions(user),
        "settings": settings_dict(),
        "products": [dict(r) | {"purchase_price": to_money(r["purchase_price_paisa"]), "selling_price": to_money(r["selling_price_paisa"])} for r in fetchall("SELECT * FROM products ORDER BY id")],
        "customers": [dict(r) for r in fetchall("SELECT id,name,phone,vehicle_no,customer_type,status FROM customers WHERE status='active' ORDER BY name")],
        "suppliers": [dict(r) for r in fetchall("SELECT id,name,phone,company,status FROM suppliers WHERE status='active' ORDER BY name")],
        "employees": [dict(r) for r in fetchall("SELECT id,name,designation,status FROM employees WHERE status='active' ORDER BY name")],
        "pumps": [dict(r) for r in fetchall("SELECT id,pump_number,active FROM pumps ORDER BY pump_number")],
        "nozzles": [dict(r) for r in fetchall("SELECT n.id,n.pump_id,n.nozzle_number,n.product_id,n.tank_id,n.active,n.last_meter,p.pump_number,pr.name AS product_name,t.tank_number FROM nozzles n JOIN pumps p ON p.id=n.pump_id JOIN products pr ON pr.id=n.product_id LEFT JOIN tanks t ON t.id=n.tank_id ORDER BY p.pump_number,n.nozzle_number")],
        "shifts": [dict(r) for r in fetchall("SELECT * FROM shifts ORDER BY id")],
        "tanks": [dict(r) | {"book_stock": tank_book_quantity(r["id"])} for r in fetchall("SELECT t.*,p.name AS product_name,p.unit FROM tanks t JOIN products p ON p.id=t.product_id ORDER BY tank_number")],
        "banks": [dict(r) | {"balance": to_money(bank_balance(r["id"]))} for r in fetchall("SELECT * FROM bank_accounts WHERE status='active' ORDER BY bank_name")],
        "expense_categories": EXPENSE_CATEGORIES,
        "customer_types": CUSTOMER_TYPES,
        "payment_methods": PAYMENT_METHODS,
        "role_labels": ROLE_LABELS,
        "permission_labels": PERMISSION_LABELS,
        "today": today_text(),
        "time": now_dt().strftime("%H:%M"),
    })


@app.get("/api/settings")
@permission_required("settings:manage")
def get_settings():
    return jsonify(settings_dict())


@app.put("/api/settings")
@permission_required("settings:manage")
def update_settings():
    data = request_data()
    allowed = set(DEFAULT_SETTINGS) - {"last_auto_backup", "purchase_next", "memo_next"}
    text_keys = {"business_name", "dealer_name", "business_address", "country", "currency", "timezone", "language", "contact_phone", "contact_mobile", "memo_prefix", "memo_format", "purchase_prefix", "logo_data", "backup_auto"}
    with atomic() as conn:
        before = settings_dict()
        changed = {}
        for key, value in data.items():
            if key not in allowed:
                continue
            if key == "logo_data" and len(str(value or "")) > 1_500_000:
                raise ApiError("The logo image is too large. Please use an image under 1 MB.")
            if key in text_keys:
                normalized = str(value or "").strip()
                if key in {"business_name", "dealer_name"} and not normalized:
                    raise ApiError(f"{key.replace('_', ' ').title()} cannot be empty.")
            elif key in {"memo_start", "opening_cash", "due_warning"}:
                if key == "memo_start":
                    number = parse_int(value, "Memo starting number")
                    if number is None or number < 1:
                        raise ApiError("Memo starting number must be at least 1.")
                    normalized = str(number)
                    existing = fetchone("SELECT memo_no FROM sales ORDER BY id DESC LIMIT 1")
                    if existing and normalized != before.get("memo_start"):
                        raise ApiError("Memo starting number cannot be changed after a memo has been issued.")
                    set_setting(conn, "memo_next", normalized)
                else:
                    minor = money_minor(value, key)
                    normalized = f"{to_money(minor):.2f}"
            else:
                normalized = str(value)
            set_setting(conn, key, normalized)
            changed[key] = normalized
        if changed:
            audit(conn, "update", "settings", "business", before, changed, data.get("reason"))
    return jsonify(settings_dict())


@app.get("/api/roles")
@permission_required("users:manage")
def list_roles():
    roles = []
    for role in fetchall("SELECT name,label FROM roles ORDER BY CASE name WHEN 'owner' THEN 1 WHEN 'manager' THEN 2 WHEN 'accountant' THEN 3 WHEN 'computer_operator' THEN 4 ELSE 5 END"):
        roles.append({"name": role["name"], "label": role["label"], "permissions": [r["permission"] for r in fetchall("SELECT permission FROM role_permissions WHERE role=? ORDER BY permission", (role["name"],))]})
    return jsonify({"roles": roles, "permission_labels": PERMISSION_LABELS})


@app.put("/api/roles/<role_name>")
@permission_required("users:manage")
def update_role_permissions(role_name: str):
    if role_name == "owner":
        raise ApiError("Owner permissions cannot be restricted.")
    data = request_data()
    permissions = data.get("permissions")
    if not isinstance(permissions, list):
        raise ApiError("Permissions must be a list.")
    if not fetchone("SELECT 1 FROM roles WHERE name=?", (role_name,)):
        raise ApiError("Role not found.", 404)
    invalid = set(permissions) - set(PERMISSION_LABELS)
    if invalid:
        raise ApiError("The request contains an unknown permission.")
    with atomic() as conn:
        previous = [r["permission"] for r in fetchall("SELECT permission FROM role_permissions WHERE role=?", (role_name,), conn)]
        conn.execute("DELETE FROM role_permissions WHERE role=?", (role_name,))
        conn.executemany("INSERT INTO role_permissions(role,permission) VALUES(?,?)", [(role_name, p) for p in sorted(set(permissions))])
        audit(conn, "permissions_updated", "users", role_name, previous, sorted(set(permissions)), data.get("reason"))
    return jsonify({"role": role_name, "permissions": sorted(set(permissions))})


@app.get("/api/users")
@permission_required("users:manage")
def list_users():
    rows = fetchall("SELECT id,username,full_name,role,is_active,created_at,last_login FROM users ORDER BY id")
    return jsonify([dict(r) for r in rows])


@app.post("/api/users")
@permission_required("users:manage")
def create_user():
    data = request_data()
    username = str(data.get("username", "")).strip()
    full_name = str(data.get("full_name", "")).strip()
    password = str(data.get("password", ""))
    role = str(data.get("role", ""))
    if len(username) < 3 or not full_name:
        raise ApiError("Enter a username (at least 3 characters) and full name.")
    if len(password) < 10:
        raise ApiError("Password must contain at least 10 characters.")
    if not fetchone("SELECT 1 FROM roles WHERE name=?", (role,)):
        raise ApiError("Choose a valid role.")
    with atomic() as conn:
        cur = conn.execute("INSERT INTO users(username,full_name,password_hash,role,is_active,created_at) VALUES(?,?,?,?,1,?)",
                           (username, full_name, generate_password_hash(password, method="scrypt"), role, now_text()))
        audit(conn, "created", "users", cur.lastrowid, None, {"username": username, "full_name": full_name, "role": role})
    return jsonify({"id": cur.lastrowid, "username": username, "full_name": full_name, "role": role, "is_active": 1}), 201


@app.put("/api/users/<int:user_id>")
@permission_required("users:manage")
def update_user(user_id: int):
    data = request_data()
    old = fetchone("SELECT id,username,full_name,role,is_active FROM users WHERE id=?", (user_id,))
    if not old:
        raise ApiError("User not found.", 404)
    if user_id == g.user["id"] and not data.get("is_active", old["is_active"]):
        raise ApiError("You cannot disable your own account.")
    full_name = str(data.get("full_name", old["full_name"])).strip()
    role = str(data.get("role", old["role"]))
    active = int(bool(data.get("is_active", old["is_active"])))
    if not full_name or not fetchone("SELECT 1 FROM roles WHERE name=?", (role,)):
        raise ApiError("Enter a name and valid role.")
    with atomic() as conn:
        conn.execute("UPDATE users SET full_name=?,role=?,is_active=? WHERE id=?", (full_name, role, active, user_id))
        if data.get("password"):
            if len(str(data["password"])) < 10:
                raise ApiError("Password must contain at least 10 characters.")
            conn.execute("UPDATE users SET password_hash=? WHERE id=?", (generate_password_hash(str(data["password"]), method="scrypt"), user_id))
        audit(conn, "updated", "users", user_id, dict(old), {"full_name": full_name, "role": role, "is_active": active, "password_changed": bool(data.get("password"))}, data.get("reason"))
    return jsonify({"ok": True})


@app.get("/api/health")
def health():
    return jsonify({"status": "ok", "database": "sqlite", "timezone": TIMEZONE_NAME})

# ---------- Products, prices and operational setup ----------
@app.get("/api/products")
@permission_required("products:view")
def list_products():
    rows = fetchall("SELECT * FROM products ORDER BY id")
    return jsonify([dict(r) | {"purchase_price": to_money(r["purchase_price_paisa"]), "selling_price": to_money(r["selling_price_paisa"])} for r in rows])


@app.post("/api/products")
@permission_required("products:manage")
def create_product():
    data = request_data()
    name = str(data.get("name", "")).strip()
    category = str(data.get("category", "other")).strip().lower()
    unit = str(data.get("unit", "পিস")).strip()
    if not name or not unit:
        raise ApiError("Product name and unit are required.")
    if category not in {"fuel", "lubricant", "other"}:
        raise ApiError("Choose a valid product category.")
    selling = money_minor(data.get("selling_price", 0), "selling price")
    purchase = money_minor(data.get("purchase_price", 0), "purchase price")
    min_stock = quantity(data.get("min_stock", 0), "minimum stock")
    stamp = now_text()
    with atomic() as conn:
        cur = conn.execute("INSERT INTO products(name,category,unit,purchase_price_paisa,selling_price_paisa,current_stock,min_stock,active,created_at,updated_at) VALUES(?,?,?,?,?,0,?,1,?,?)",
                           (name, category, unit, purchase, selling, min_stock, stamp, stamp))
        if selling:
            conn.execute("INSERT INTO price_history(product_id,old_price_paisa,new_price_paisa,effective_at,user_id,reason) VALUES(?,0,?,?,?,?)",
                         (cur.lastrowid, selling, stamp, g.user["id"], "Initial selling price"))
        audit(conn, "created", "products", cur.lastrowid, None, {"name": name, "category": category, "unit": unit, "selling_price": to_money(selling)})
    return jsonify({"id": cur.lastrowid, "name": name}), 201


@app.put("/api/products/<int:product_id>")
@permission_required("products:manage")
def update_product(product_id: int):
    data = request_data()
    old = active_record("products", product_id)
    name = str(data.get("name", old["name"])).strip()
    category = str(data.get("category", old["category"])).strip().lower()
    unit = str(data.get("unit", old["unit"])).strip()
    if not name or not unit or category not in {"fuel", "lubricant", "other"}:
        raise ApiError("Enter a product name, valid category and unit.")
    new_price = money_minor(data.get("selling_price", to_money(old["selling_price_paisa"])), "selling price")
    new_cost = money_minor(data.get("purchase_price", to_money(old["purchase_price_paisa"])), "purchase price")
    minimum = quantity(data.get("min_stock", old["min_stock"]), "minimum stock")
    active = int(bool(data.get("active", old["active"])))
    before = {"name": old["name"], "category": old["category"], "unit": old["unit"], "selling_price": to_money(old["selling_price_paisa"]), "purchase_price": to_money(old["purchase_price_paisa"]), "min_stock": old["min_stock"], "active": old["active"]}
    after = {"name": name, "category": category, "unit": unit, "selling_price": to_money(new_price), "purchase_price": to_money(new_cost), "min_stock": minimum, "active": active}
    if new_price != old["selling_price_paisa"] and "prices:manage" not in user_permissions(g.user):
        raise ApiError("You do not have permission to change selling prices.", 403, "permission_denied")
    with atomic() as conn:
        conn.execute("UPDATE products SET name=?,category=?,unit=?,purchase_price_paisa=?,selling_price_paisa=?,min_stock=?,active=?,updated_at=? WHERE id=?",
                     (name, category, unit, new_cost, new_price, minimum, active, now_text(), product_id))
        if new_price != old["selling_price_paisa"]:
            conn.execute("INSERT INTO price_history(product_id,old_price_paisa,new_price_paisa,effective_at,user_id,reason) VALUES(?,?,?,?,?,?)",
                         (product_id, old["selling_price_paisa"], new_price, now_text(), g.user["id"], str(data.get("reason", ""))[:500]))
            audit(conn, "price_changed", "prices", product_id, {"price": to_money(old["selling_price_paisa"])}, {"price": to_money(new_price)}, data.get("reason"))
        audit(conn, "updated", "products", product_id, before, after, data.get("reason"))
    return jsonify({"ok": True})


@app.get("/api/price-history")
@permission_required("products:view")
def price_history():
    rows = fetchall("SELECT h.*,p.name AS product_name,u.full_name AS user_name FROM price_history h JOIN products p ON p.id=h.product_id LEFT JOIN users u ON u.id=h.user_id ORDER BY h.id DESC LIMIT 1000")
    return jsonify([dict(r) | {"old_price": to_money(r["old_price_paisa"]), "new_price": to_money(r["new_price_paisa"])} for r in rows])


@app.post("/api/stock/adjustments")
@permission_required("inventory:adjust")
def stock_adjustment():
    data = request_data()
    product_id = parse_int(data.get("product_id"), "product")
    product = active_record("products", product_id)
    change = quantity(data.get("quantity"), "adjustment quantity", allow_negative=True)
    if not change:
        raise ApiError("Adjustment quantity cannot be zero.")
    reason = str(data.get("reason", "")).strip()
    if not reason:
        raise ApiError("A reason is required for every stock adjustment.")
    unit_cost = money_minor(data.get("unit_cost", to_money(product["purchase_price_paisa"])), "unit cost")
    with atomic() as conn:
        old_stock = float(product["current_stock"])
        updated = old_stock + change
        if updated < -0.0005:
            raise ApiError("This adjustment would make book stock negative.")
        new_cost = int(product["purchase_price_paisa"])
        if change > 0 and unit_cost > 0:
            if old_stock > 0:
                new_cost = int(((Decimal(str(old_stock)) * Decimal(new_cost) + Decimal(str(change)) * Decimal(unit_cost)) / Decimal(str(updated))).quantize(Decimal("1"), rounding=ROUND_HALF_UP))
            else:
                new_cost = unit_cost
        conn.execute("UPDATE products SET current_stock=?,purchase_price_paisa=?,updated_at=? WHERE id=?", (round(updated, 3), new_cost, now_text(), product_id))
        cur = conn.execute("INSERT INTO stock_movements(product_id,kind,quantity,business_date,unit_cost_paisa,reference_type,notes,user_id,created_at) VALUES(?,'adjustment',?,?,?,'adjustment',?,?,?)",
                           (product_id, change, today_text(), unit_cost if change > 0 else product["purchase_price_paisa"], reason, g.user["id"], now_text()))
        audit(conn, "stock_adjusted", "inventory", product_id, {"stock": old_stock, "average_cost": to_money(product["purchase_price_paisa"])}, {"stock": round(updated, 3), "adjustment": change, "average_cost": to_money(new_cost)}, reason)
    return jsonify({"id": cur.lastrowid, "current_stock": round(updated, 3), "average_cost": to_money(new_cost)}), 201


@app.get("/api/stock/movements")
@permission_required("inventory:view")
def stock_movements():
    query = str(request.args.get("q", "")).strip()
    params: list[Any] = []
    sql = "SELECT m.*,p.name AS product_name,p.unit,u.full_name AS user_name FROM stock_movements m JOIN products p ON p.id=m.product_id LEFT JOIN users u ON u.id=m.user_id WHERE 1=1"
    if request.args.get("product_id"):
        sql += " AND m.product_id=?"; params.append(parse_int(request.args["product_id"], "product"))
    if query:
        sql += " AND (p.name LIKE ? OR m.notes LIKE ? OR m.reference_type LIKE ?)"; params.extend([f"%{query}%"] * 3)
    sql += " ORDER BY m.id DESC LIMIT 1000"
    return jsonify([dict(r) for r in fetchall(sql, tuple(params))])


@app.get("/api/attendance")
@permission_required("employees:manage")
def list_attendance():
    start = parse_date(request.args.get("start"), "start date") if request.args.get("start") else "0001-01-01"
    end = parse_date(request.args.get("end"), "end date") if request.args.get("end") else today_text()
    rows = fetchall("SELECT a.*,e.name AS employee_name,e.designation,s.name AS shift_name,u.full_name AS recorded_by FROM attendance a JOIN employees e ON e.id=a.employee_id JOIN shifts s ON s.id=a.shift_id LEFT JOIN users u ON u.id=a.created_by WHERE a.date BETWEEN ? AND ? ORDER BY a.date DESC,a.id DESC LIMIT 2000", (start,end))
    return jsonify([dict(r) for r in rows])


@app.post("/api/attendance")
@permission_required("employees:manage")
def record_attendance():
    data = request_data()
    employee_id = parse_int(data.get("employee_id"), "employee")
    shift_id = parse_int(data.get("shift_id"), "shift")
    employee = active_record("employees", employee_id)
    shift = active_record("shifts", shift_id)
    day = parse_date(data.get("date"), "date")
    status = str(data.get("status", "present"))
    if status not in {"present", "absent", "leave"}:
        raise ApiError("Attendance status must be present, absent or leave.")
    check_in = str(data.get("check_in", "")).strip() or None
    check_out = str(data.get("check_out", "")).strip() or None
    for value in (check_in, check_out):
        if value:
            try: datetime.strptime(value, "%H:%M")
            except ValueError: raise ApiError("Attendance times must use HH:MM format.")
    old = fetchone("SELECT * FROM attendance WHERE employee_id=? AND date=? AND shift_id=?", (employee_id,day,shift_id))
    with atomic() as conn:
        if old:
            conn.execute("UPDATE attendance SET status=?,check_in=?,check_out=?,notes=?,created_by=?,updated_at=? WHERE id=?", (status,check_in,check_out,str(data.get("notes","")),g.user["id"],now_text(),old["id"]))
            record_id=old["id"]
            audit(conn,"attendance_updated","employees",record_id,dict(old),{"employee":employee["name"],"shift":shift["name"],"date":day,"status":status,"check_in":check_in,"check_out":check_out})
        else:
            cur=conn.execute("INSERT INTO attendance(employee_id,date,shift_id,status,check_in,check_out,notes,created_by,created_at,updated_at) VALUES(?,?,?,?,?,?,?,?,?,?)",(employee_id,day,shift_id,status,check_in,check_out,str(data.get("notes","")),g.user["id"],now_text(),now_text()))
            record_id=cur.lastrowid
            audit(conn,"attendance_recorded","employees",record_id,None,{"employee":employee["name"],"shift":shift["name"],"date":day,"status":status,"check_in":check_in,"check_out":check_out})
    return jsonify({"id":record_id,"updated":bool(old)}),201


@app.get("/api/pumps")
@permission_required("products:view")
def list_pumps():
    pumps = []
    for p in fetchall("SELECT * FROM pumps ORDER BY pump_number"):
        pitem = dict(p)
        pitem["nozzles"] = [dict(n) for n in fetchall("SELECT n.*,pr.name AS product_name,t.tank_number FROM nozzles n JOIN products pr ON pr.id=n.product_id LEFT JOIN tanks t ON t.id=n.tank_id WHERE n.pump_id=? ORDER BY n.nozzle_number", (p["id"],))]
        pumps.append(pitem)
    return jsonify(pumps)


@app.post("/api/pumps")
@permission_required("pumps:manage")
def create_pump():
    data = request_data()
    number = str(data.get("pump_number", "")).strip()
    if not number:
        raise ApiError("Pump number is required.")
    with atomic() as conn:
        cur = conn.execute("INSERT INTO pumps(pump_number,active,created_at) VALUES(?,?,?)", (number, int(bool(data.get("active", True))), now_text()))
        audit(conn, "created", "pumps", cur.lastrowid, None, {"pump_number": number})
    return jsonify({"id": cur.lastrowid, "pump_number": number}), 201


@app.put("/api/pumps/<int:pump_id>")
@permission_required("pumps:manage")
def update_pump(pump_id: int):
    old = active_record("pumps", pump_id)
    data = request_data()
    number = str(data.get("pump_number", old["pump_number"])).strip()
    active = int(bool(data.get("active", old["active"])))
    with atomic() as conn:
        conn.execute("UPDATE pumps SET pump_number=?,active=? WHERE id=?", (number, active, pump_id))
        audit(conn, "updated", "pumps", pump_id, dict(old), {"pump_number": number, "active": active})
    return jsonify({"ok": True})


@app.post("/api/nozzles")
@permission_required("pumps:manage")
def create_nozzle():
    data = request_data()
    pump_id = parse_int(data.get("pump_id"), "pump")
    product_id = parse_int(data.get("product_id"), "product")
    number = str(data.get("nozzle_number", "")).strip()
    if not number:
        raise ApiError("Nozzle number is required.")
    active_record("pumps", pump_id)
    product = active_record("products", product_id)
    tank_id = parse_int(data.get("tank_id"), "tank", None)
    if tank_id:
        tank = active_record("tanks", tank_id)
        if tank["product_id"] != product_id:
            raise ApiError("The selected tank is assigned to a different product.")
    with atomic() as conn:
        cur = conn.execute("INSERT INTO nozzles(pump_id,nozzle_number,product_id,tank_id,active) VALUES(?,?,?,?,?)", (pump_id, number, product_id, tank_id, int(bool(data.get("active", True)))))
        audit(conn, "created", "nozzles", cur.lastrowid, None, {"pump_id": pump_id, "nozzle_number": number, "product_id": product_id, "tank_id": tank_id})
    return jsonify({"id": cur.lastrowid}), 201


@app.put("/api/nozzles/<int:nozzle_id>")
@permission_required("pumps:manage")
def update_nozzle(nozzle_id: int):
    old = active_record("nozzles", nozzle_id)
    data = request_data()
    product_id = parse_int(data.get("product_id", old["product_id"]), "product")
    active_record("products", product_id)
    tank_id = parse_int(data.get("tank_id", old["tank_id"]), "tank", None)
    if tank_id:
        tank = active_record("tanks", tank_id)
        if tank["product_id"] != product_id:
            raise ApiError("The selected tank is assigned to a different product.")
    number = str(data.get("nozzle_number", old["nozzle_number"])).strip()
    active = int(bool(data.get("active", old["active"])))
    with atomic() as conn:
        conn.execute("UPDATE nozzles SET nozzle_number=?,product_id=?,tank_id=?,active=? WHERE id=?", (number, product_id, tank_id, active, nozzle_id))
        audit(conn, "updated", "nozzles", nozzle_id, dict(old), {"nozzle_number": number, "product_id": product_id, "tank_id": tank_id, "active": active})
    return jsonify({"ok": True})


@app.get("/api/shifts")
@permission_required("shifts:manage")
def list_shifts():
    rows = fetchall("SELECT s.*, (SELECT COUNT(*) FROM shift_closings c WHERE c.shift_id=s.id) AS closing_count FROM shifts s ORDER BY id")
    return jsonify([dict(r) for r in rows])


@app.post("/api/shifts")
@permission_required("shifts:manage")
def create_shift():
    data = request_data()
    name = str(data.get("name", "")).strip()
    start, end = str(data.get("start_time", "")).strip(), str(data.get("end_time", "")).strip()
    for value in (start, end):
        try: datetime.strptime(value, "%H:%M")
        except ValueError: raise ApiError("Shift times must use 24-hour HH:MM format.")
    if not name:
        raise ApiError("Shift name is required.")
    with atomic() as conn:
        cur = conn.execute("INSERT INTO shifts(name,start_time,end_time,active,created_at) VALUES(?,?,?,?,?)", (name, start, end, int(bool(data.get("active", True))), now_text()))
        audit(conn, "created", "shifts", cur.lastrowid, None, {"name": name, "start_time": start, "end_time": end})
    return jsonify({"id": cur.lastrowid}), 201


@app.put("/api/shifts/<int:shift_id>")
@permission_required("shifts:manage")
def update_shift(shift_id: int):
    old = active_record("shifts", shift_id)
    data = request_data()
    name = str(data.get("name", old["name"])).strip()
    start, end = str(data.get("start_time", old["start_time"])), str(data.get("end_time", old["end_time"]))
    for value in (start, end):
        try: datetime.strptime(value, "%H:%M")
        except ValueError: raise ApiError("Shift times must use 24-hour HH:MM format.")
    active = int(bool(data.get("active", old["active"])))
    with atomic() as conn:
        conn.execute("UPDATE shifts SET name=?,start_time=?,end_time=?,active=? WHERE id=?", (name, start, end, active, shift_id))
        audit(conn, "updated", "shifts", shift_id, dict(old), {"name": name, "start_time": start, "end_time": end, "active": active})
    return jsonify({"ok": True})


@app.post("/api/shifts/<int:shift_id>/close")
@permission_required("shifts:manage")
def close_shift(shift_id: int):
    shift = active_record("shifts", shift_id)
    data = request_data()
    day = parse_date(data.get("date"), "date")
    opening = money_minor(data.get("opening_cash", 0), "opening cash")
    closing = money_minor(data.get("closing_cash", 0), "closing cash")
    sales = fetchone("SELECT COALESCE(SUM(total_paisa),0) total,COALESCE(SUM(paid_cash_paisa),0) cash,COALESCE(SUM(credit_paisa),0) credit FROM sales WHERE date=? AND shift_id=? AND status='posted'", (day, shift_id))
    expenses = fetchone("SELECT COALESCE(SUM(amount_paisa),0) amount FROM expenses WHERE date=? AND shift_id=? AND status='approved'", (day, shift_id))
    deposits = fetchone("SELECT COALESCE(SUM(amount_paisa),0) amount FROM bank_transactions WHERE date=? AND shift_id=? AND type='cash_deposit' AND direction='in'", (day, shift_id))
    cash_flows = fetchone("SELECT COALESCE(SUM(CASE WHEN direction='in' THEN amount_paisa ELSE 0 END),0) cash_in,COALESCE(SUM(CASE WHEN direction='out' THEN amount_paisa ELSE 0 END),0) cash_out FROM cash_transactions WHERE date=? AND shift_id=?", (day, shift_id))
    meters = fetchone("SELECT COALESCE(MIN(opening_meter),0) opening,COALESCE(MAX(closing_meter),0) closing,COALESCE(SUM(sales_liters),0) liters FROM meter_readings WHERE date=? AND shift_id=?", (day, shift_id))
    cash_sales, credit_sales = int(sales["cash"] or 0), int(sales["credit"] or 0)
    expected = opening + int(cash_flows["cash_in"] or 0) - int(cash_flows["cash_out"] or 0)
    difference = closing - expected
    with atomic() as conn:
        cur = conn.execute("INSERT INTO shift_closings(date,shift_id,opening_cash_paisa,closing_cash_paisa,opening_meter,closing_meter,total_liters,total_sales_paisa,cash_sales_paisa,credit_sales_paisa,expenses_paisa,deposits_paisa,expected_cash_paisa,cash_difference_paisa,notes,closed_by,closed_at) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
                           (day, shift_id, opening, closing, meters["opening"], meters["closing"], meters["liters"], sales["total"], cash_sales, credit_sales, int(expenses["amount"] or 0), int(deposits["amount"] or 0), expected, difference, str(data.get("notes", "")), g.user["id"], now_text()))
        audit(conn, "shift_closed", "shifts", cur.lastrowid, None, {"date": day, "shift": shift["name"], "opening_meter": meters["opening"], "closing_meter": meters["closing"], "total_liters": meters["liters"], "total_sales": to_money(sales["total"]), "expected_cash": to_money(expected), "closing_cash": to_money(closing), "cash_difference": to_money(difference)})
    return jsonify({"id": cur.lastrowid, "expected_cash": to_money(expected), "cash_difference": to_money(difference), "opening_meter": meters["opening"], "closing_meter": meters["closing"], "total_liters": meters["liters"], "total_sales": to_money(sales["total"]) }), 201


@app.get("/api/shift-closings")
@permission_required("reports:view")
def list_shift_closings():
    start, end = date_range_params()
    rows = fetchall("SELECT c.*,s.name AS shift_name,u.full_name AS closed_by_name FROM shift_closings c JOIN shifts s ON s.id=c.shift_id LEFT JOIN users u ON u.id=c.closed_by WHERE c.date BETWEEN ? AND ? ORDER BY c.date DESC,c.id DESC", (start, end))
    money_fields = ("opening_cash_paisa", "closing_cash_paisa", "total_sales_paisa", "cash_sales_paisa", "credit_sales_paisa", "expenses_paisa", "deposits_paisa", "expected_cash_paisa", "cash_difference_paisa")
    return jsonify([{**dict(r), **{f.removesuffix("_paisa"): to_money(r[f]) for f in money_fields}} for r in rows])


@app.get("/api/tanks")
@permission_required("inventory:view")
def list_tanks():
    rows = fetchall("SELECT t.*,p.name AS product_name,p.unit,p.current_stock FROM tanks t JOIN products p ON p.id=t.product_id ORDER BY tank_number")
    return jsonify([dict(r) | {"book_stock": tank_book_quantity(r["id"])} for r in rows])


@app.post("/api/tanks")
@permission_required("tanks:manage")
def create_tank():
    data = request_data()
    number = str(data.get("tank_number", "")).strip()
    product_id = parse_int(data.get("product_id"), "product")
    product = active_record("products", product_id)
    capacity = quantity(data.get("capacity"), "capacity")
    opening = quantity(data.get("opening_quantity", 0), "opening quantity")
    minimum = quantity(data.get("min_level", 0), "minimum level")
    maximum = quantity(data.get("max_level", capacity), "maximum level")
    if not number or capacity <= 0 or opening > capacity or maximum > capacity:
        raise ApiError("Enter a tank number and valid capacity / stock levels.")
    stamp = now_text()
    with atomic() as conn:
        cur = conn.execute("INSERT INTO tanks(tank_number,product_id,capacity,opening_quantity,min_level,max_level,active,created_at) VALUES(?,?,?,?,?,?,1,?)", (number, product_id, capacity, opening, minimum, maximum, stamp))
        audit(conn, "created", "tanks", cur.lastrowid, None, {"tank_number": number, "product": product["name"], "capacity": capacity})
    return jsonify({"id": cur.lastrowid}), 201


@app.put("/api/tanks/<int:tank_id>")
@permission_required("tanks:manage")
def update_tank(tank_id: int):
    old = active_record("tanks", tank_id)
    data = request_data()
    product_id = parse_int(data.get("product_id", old["product_id"]), "product")
    active_record("products", product_id)
    number = str(data.get("tank_number", old["tank_number"])).strip()
    capacity = quantity(data.get("capacity", old["capacity"]), "capacity")
    opening = quantity(data.get("opening_quantity", old["opening_quantity"]), "opening quantity")
    minimum = quantity(data.get("min_level", old["min_level"]), "minimum level")
    maximum = quantity(data.get("max_level", old["max_level"]), "maximum level")
    active = int(bool(data.get("active", old["active"])))
    if not number or capacity <= 0 or opening > capacity or maximum > capacity:
        raise ApiError("Enter valid tank details; quantity and maximum level cannot exceed capacity.")
    with atomic() as conn:
        conn.execute("UPDATE tanks SET tank_number=?,product_id=?,capacity=?,opening_quantity=?,min_level=?,max_level=?,active=? WHERE id=?", (number, product_id, capacity, opening, minimum, maximum, active, tank_id))
        audit(conn, "updated", "tanks", tank_id, dict(old), {"tank_number": number, "product_id": product_id, "capacity": capacity, "opening_quantity": opening, "min_level": minimum, "max_level": maximum, "active": active})
    return jsonify({"ok": True})


@app.post("/api/tanks/<int:tank_id>/dips")
@permission_required("tanks:manage")
def create_tank_dip(tank_id: int):
    tank = active_record("tanks", tank_id)
    data = request_data()
    day = parse_date(data.get("date"), "date")
    dip = quantity(data.get("dip_reading"), "dip reading")
    estimated = quantity(data.get("estimated_quantity"), "estimated quantity")
    book = tank_book_quantity(tank_id)
    physical = quantity(data.get("physical_stock", estimated), "physical stock")
    if dip > tank["capacity"] or estimated > tank["capacity"] or physical > tank["capacity"]:
        raise ApiError("Dip or estimated stock exceeds tank capacity.")
    difference = round(physical - book, 3)
    operator_id = parse_int(data.get("operator_id"), "operator", None)
    if operator_id:
        active_record("employees", operator_id)
    with atomic() as conn:
        cur = conn.execute("INSERT INTO tank_dips(date,tank_id,dip_reading,estimated_quantity,book_stock,physical_stock,difference,operator_id,verified,notes,user_id,created_at) VALUES(?,?,?,?,?,?,?,?,0,?,?,?)",
                           (day, tank_id, dip, estimated, book, physical, difference, operator_id, str(data.get("notes", "")), g.user["id"], now_text()))
        audit(conn, "dip_recorded", "tanks", cur.lastrowid, None, {"tank": tank["tank_number"], "book_stock": book, "physical_stock": physical, "difference": difference}, data.get("notes"))
    return jsonify({"id": cur.lastrowid, "verified": False, "difference": difference}), 201


@app.post("/api/tank-dips/<int:dip_id>/verify")
@permission_required("tanks:manage")
def verify_tank_dip(dip_id: int):
    dip = fetchone("SELECT * FROM tank_dips WHERE id=?", (dip_id,))
    if not dip:
        raise ApiError("Dip reading not found.", 404)
    with atomic() as conn:
        conn.execute("UPDATE tank_dips SET verified=1,verified_by=? WHERE id=?", (g.user["id"], dip_id))
        audit(conn, "dip_verified", "tanks", dip_id, {"verified": bool(dip["verified"])}, {"verified": True}, request_data().get("reason"))
    return jsonify({"ok": True})

# ---------- Customers, suppliers, employees and balances ----------
@app.get("/api/customers")
@permission_required("customers:manage")
def list_customers():
    query = str(request.args.get("q", "")).strip()
    sql = """SELECT c.*, c.opening_due_paisa + COALESCE((SELECT SUM(t.amount_paisa) FROM customer_transactions t WHERE t.customer_id=c.id),0) AS current_due_paisa,
        (SELECT MAX(date) FROM customer_transactions t WHERE t.customer_id=c.id AND t.type IN ('credit_sale','sale_edit_credit')) AS last_purchase,
        (SELECT MAX(date) FROM customer_transactions t WHERE t.customer_id=c.id AND t.type='payment') AS last_payment
        FROM customers c WHERE 1=1"""
    params: list[Any] = []
    if query:
        sql += " AND (c.name LIKE ? OR c.phone LIKE ? OR c.vehicle_no LIKE ?)"
        params.extend([f"%{query}%"] * 3)
    sql += " ORDER BY c.name LIMIT 1000"
    return jsonify([dict(r) | {"opening_due": to_money(r["opening_due_paisa"]), "current_due": to_money(r["current_due_paisa"])} for r in fetchall(sql, tuple(params))])


@app.post("/api/customers")
@permission_required("customers:manage")
def create_customer():
    data = request_data()
    name = str(data.get("name", "")).strip()
    if not name:
        raise ApiError("Customer name is required.")
    kind = str(data.get("customer_type", "Regular Customer"))
    if kind not in CUSTOMER_TYPES:
        raise ApiError("Choose a valid customer type.")
    opening = money_minor(data.get("opening_due", 0), "opening due")
    stamp = now_text()
    with atomic() as conn:
        cur = conn.execute("INSERT INTO customers(name,phone,address,vehicle_no,customer_type,opening_due_paisa,status,created_at,updated_at) VALUES(?,?,?,?,?,?, 'active',?,?)",
                           (name, str(data.get("phone", "")).strip(), str(data.get("address", "")).strip(), str(data.get("vehicle_no", "")).strip(), kind, opening, stamp, stamp))
        audit(conn, "created", "customers", cur.lastrowid, None, {"name": name, "phone": data.get("phone"), "vehicle_no": data.get("vehicle_no"), "opening_due": to_money(opening)})
    return jsonify({"id": cur.lastrowid, "name": name}), 201


@app.put("/api/customers/<int:customer_id>")
@permission_required("customers:manage")
def update_customer(customer_id: int):
    old = active_record("customers", customer_id)
    data = request_data()
    name = str(data.get("name", old["name"])).strip()
    kind = str(data.get("customer_type", old["customer_type"]))
    if not name or kind not in CUSTOMER_TYPES:
        raise ApiError("Enter a name and valid customer type.")
    # Existing due is only changed through an audited adjustment, not by editing the profile.
    active = str(data.get("status", old["status"]))
    if active not in {"active", "inactive"}:
        raise ApiError("Customer status must be active or inactive.")
    after = {"name": name, "phone": str(data.get("phone", old["phone"] or "")).strip(), "address": str(data.get("address", old["address"] or "")).strip(),
             "vehicle_no": str(data.get("vehicle_no", old["vehicle_no"] or "")).strip(), "customer_type": kind, "status": active}
    with atomic() as conn:
        conn.execute("UPDATE customers SET name=?,phone=?,address=?,vehicle_no=?,customer_type=?,status=?,updated_at=? WHERE id=?",
                     (after["name"], after["phone"], after["address"], after["vehicle_no"], kind, active, now_text(), customer_id))
        audit(conn, "updated", "customers", customer_id, {"name": old["name"], "phone": old["phone"], "vehicle_no": old["vehicle_no"], "status": old["status"]}, after)
    return jsonify({"ok": True})


@app.post("/api/customers/<int:customer_id>/payments")
@permission_required("due:collect")
def customer_payment(customer_id: int):
    customer = active_record("customers", customer_id)
    data = request_data()
    amount = money_minor(data.get("amount"), "payment")
    ensure_positive(amount, "Payment")
    day = parse_date(data.get("date"), "date")
    shift_id = parse_int(data.get("shift_id"), "shift", None)
    if shift_id: active_record("shifts", shift_id)
    method = str(data.get("payment_method", "cash"))
    if method not in {"cash", "bank", "mobile"}:
        raise ApiError("Select cash, bank or mobile banking.")
    account_id = parse_int(data.get("bank_account_id"), "bank account", None)
    if method == "bank" and not account_id:
        raise ApiError("Select the bank account where this payment was received.")
    if account_id:
        active_record("bank_accounts", account_id)
    current_due = balance_customer(customer_id)
    if amount > current_due:
        raise ApiError(f"Payment exceeds the customer's current due of ৳{to_money(current_due):,.2f}.")
    with atomic() as conn:
        cur = conn.execute("INSERT INTO customer_transactions(customer_id,date,type,amount_paisa,source,notes,user_id,created_at) VALUES(?,?,'payment',?,'payment',?,?,?)",
                           (customer_id, day, -amount, str(data.get("notes", "")), g.user["id"], now_text()))
        if method == "cash":
            conn.execute("INSERT INTO cash_transactions(date,shift_id,direction,type,amount_paisa,reference_type,reference_id,notes,user_id,created_at) VALUES(?,?,'in','customer_payment',?,'customer_payment',?,?,?,?)",
                         (day, shift_id, amount, cur.lastrowid, str(data.get("notes", "")), g.user["id"], now_text()))
        else:
            conn.execute("INSERT INTO bank_transactions(date,bank_id,channel,direction,type,amount_paisa,reference,notes,reference_type,reference_id,shift_id,user_id,created_at) VALUES(?,?,?,'in','customer_payment',?,?,?,?,?,?,?,?)",
                         (day, account_id, method, amount, customer["name"], str(data.get("notes", "")), "customer_payment", cur.lastrowid, shift_id, g.user["id"], now_text()))
        new_due = current_due - amount
        audit(conn, "payment_received", "customer_due", customer_id, {"due": to_money(current_due)}, {"payment": to_money(amount), "remaining_due": to_money(new_due)}, data.get("notes"))
    return jsonify({"id": cur.lastrowid, "remaining_due": to_money(new_due)}), 201


@app.get("/api/customers/<int:customer_id>/statement")
@permission_required("reports:view")
def customer_statement(customer_id: int):
    customer = active_record("customers", customer_id)
    start = parse_date(request.args.get("start"), "start date") if request.args.get("start") else "0001-01-01"
    end = parse_date(request.args.get("end"), "end date") if request.args.get("end") else today_text()
    rows = fetchall("SELECT * FROM customer_transactions WHERE customer_id=? AND date BETWEEN ? AND ? ORDER BY date,id", (customer_id, start, end))
    balance = int(customer["opening_due_paisa"])
    result = []
    if start <= "0001-01-01":
        result.append({"date": "", "type": "Opening balance", "description": "Opening due", "debit": 0, "credit": to_money(balance), "balance": to_money(balance)})
    else:
        prior = fetchone("SELECT COALESCE(SUM(amount_paisa),0) value FROM customer_transactions WHERE customer_id=? AND date<?", (customer_id, start))
        balance += int(prior["value"] or 0)
        result.append({"date": start, "type": "Opening balance", "description": "Balance brought forward", "debit": 0, "credit": to_money(balance), "balance": to_money(balance)})
    for row in rows:
        amount = int(row["amount_paisa"])
        balance += amount
        result.append({"id": row["id"], "date": row["date"], "type": row["type"], "description": row["notes"] or row["source"] or row["type"],
                       "debit": to_money(max(-amount, 0)), "credit": to_money(max(amount, 0)), "balance": to_money(balance)})
    return jsonify({"customer": dict(customer) | {"opening_due": to_money(customer["opening_due_paisa"])}, "rows": result, "current_due": to_money(balance)})


@app.get("/api/suppliers")
@permission_required("purchases:view")
def list_suppliers():
    query = str(request.args.get("q", "")).strip()
    sql = "SELECT s.*, s.opening_balance_paisa + COALESCE((SELECT SUM(t.amount_paisa) FROM supplier_transactions t WHERE t.supplier_id=s.id),0) AS current_due_paisa FROM suppliers s"
    params: list[Any] = []
    if query:
        sql += " WHERE s.name LIKE ? OR s.company LIKE ? OR s.phone LIKE ?"; params.extend([f"%{query}%"] * 3)
    sql += " ORDER BY s.name LIMIT 1000"
    return jsonify([dict(r) | {"opening_balance": to_money(r["opening_balance_paisa"]), "current_due": to_money(r["current_due_paisa"])} for r in fetchall(sql, tuple(params))])


@app.post("/api/suppliers")
@permission_required("suppliers:manage")
def create_supplier():
    data = request_data()
    name = str(data.get("name", "")).strip()
    if not name:
        raise ApiError("Supplier name is required.")
    opening = money_minor(data.get("opening_balance", 0), "opening balance")
    stamp = now_text()
    with atomic() as conn:
        cur = conn.execute("INSERT INTO suppliers(name,phone,address,company,opening_balance_paisa,status,created_at,updated_at) VALUES(?,?,?,?,?,'active',?,?)",
                           (name, str(data.get("phone", "")).strip(), str(data.get("address", "")).strip(), str(data.get("company", "")).strip(), opening, stamp, stamp))
        audit(conn, "created", "suppliers", cur.lastrowid, None, {"name": name, "company": data.get("company"), "opening_balance": to_money(opening)})
    return jsonify({"id": cur.lastrowid}), 201


@app.put("/api/suppliers/<int:supplier_id>")
@permission_required("suppliers:manage")
def update_supplier(supplier_id: int):
    old = active_record("suppliers", supplier_id)
    data = request_data()
    name = str(data.get("name", old["name"])).strip()
    if not name:
        raise ApiError("Supplier name is required.")
    status = str(data.get("status", old["status"]))
    if status not in {"active", "inactive"}:
        raise ApiError("Choose active or inactive status.")
    new = {"name": name, "phone": str(data.get("phone", old["phone"] or "")).strip(), "address": str(data.get("address", old["address"] or "")).strip(),
           "company": str(data.get("company", old["company"] or "")).strip(), "status": status}
    with atomic() as conn:
        conn.execute("UPDATE suppliers SET name=?,phone=?,address=?,company=?,status=?,updated_at=? WHERE id=?", (new["name"], new["phone"], new["address"], new["company"], status, now_text(), supplier_id))
        audit(conn, "updated", "suppliers", supplier_id, {"name": old["name"], "phone": old["phone"], "company": old["company"], "status": old["status"]}, new)
    return jsonify({"ok": True})


@app.post("/api/suppliers/<int:supplier_id>/payments")
@permission_required("suppliers:manage")
def supplier_payment(supplier_id: int):
    supplier = active_record("suppliers", supplier_id)
    data = request_data()
    amount = money_minor(data.get("amount"), "payment")
    ensure_positive(amount, "Payment")
    day = parse_date(data.get("date"), "date")
    shift_id = parse_int(data.get("shift_id"), "shift", None)
    if shift_id: active_record("shifts", shift_id)
    method = str(data.get("payment_method", "cash"))
    account_id = parse_int(data.get("bank_account_id"), "bank account", None)
    if method not in {"cash", "bank", "mobile"}:
        raise ApiError("Select cash, bank or mobile banking.")
    if method == "bank" and not account_id:
        raise ApiError("Select the bank account used for payment.")
    if account_id:
        active_record("bank_accounts", account_id)
    balance = balance_supplier(supplier_id)
    if amount > balance:
        raise ApiError(f"Payment exceeds supplier due of ৳{to_money(balance):,.2f}.")
    with atomic() as conn:
        cur = conn.execute("INSERT INTO supplier_transactions(supplier_id,date,type,amount_paisa,source,notes,user_id,created_at) VALUES(?,?,'payment',?,'payment',?,?,?)",
                           (supplier_id, day, -amount, str(data.get("notes", "")), g.user["id"], now_text()))
        if method == "cash":
            conn.execute("INSERT INTO cash_transactions(date,shift_id,direction,type,amount_paisa,reference_type,reference_id,notes,user_id,created_at) VALUES(?,?,'out','supplier_payment',?,'supplier_payment',?,?,?,?)",
                         (day, shift_id, amount, cur.lastrowid, f"{supplier['name']}: {str(data.get('notes', ''))}".strip(), g.user["id"], now_text()))
        else:
            conn.execute("INSERT INTO bank_transactions(date,bank_id,channel,direction,type,amount_paisa,reference,notes,reference_type,reference_id,shift_id,user_id,created_at) VALUES(?,?,?,'out','supplier_payment',?,?,?,?,?,?,?,?)",
                         (day, account_id, method, amount, supplier["name"], str(data.get("notes", "")), "supplier_payment", cur.lastrowid, shift_id, g.user["id"], now_text()))
        audit(conn, "supplier_payment", "suppliers", supplier_id, {"due": to_money(balance)}, {"payment": to_money(amount), "remaining_due": to_money(balance - amount)}, data.get("notes"))
    return jsonify({"id": cur.lastrowid, "remaining_due": to_money(balance - amount)}), 201


@app.get("/api/employees")
@permission_required("employees:manage")
def list_employees():
    rows = fetchall("SELECT e.*,u.username FROM employees e LEFT JOIN users u ON u.id=e.user_id ORDER BY e.name")
    return jsonify([dict(r) | {"salary": to_money(r["salary_paisa"])} for r in rows])


@app.post("/api/employees")
@permission_required("employees:manage")
def create_employee():
    data = request_data()
    name = str(data.get("name", "")).strip()
    if not name:
        raise ApiError("Employee name is required.")
    salary = money_minor(data.get("salary", 0), "salary")
    joining = parse_date(data.get("join_date"), "joining date") if data.get("join_date") else None
    stamp = now_text()
    with atomic() as conn:
        cur = conn.execute("INSERT INTO employees(name,phone,address,designation,join_date,salary_paisa,status,created_at,updated_at) VALUES(?,?,?,?,?,?,'active',?,?)",
                           (name, str(data.get("phone", "")).strip(), str(data.get("address", "")).strip(), str(data.get("designation", "")).strip(), joining, salary, stamp, stamp))
        audit(conn, "created", "employees", cur.lastrowid, None, {"name": name, "designation": data.get("designation"), "salary": to_money(salary)})
    return jsonify({"id": cur.lastrowid}), 201


@app.put("/api/employees/<int:employee_id>")
@permission_required("employees:manage")
def update_employee(employee_id: int):
    old = active_record("employees", employee_id)
    data = request_data()
    name = str(data.get("name", old["name"])).strip()
    if not name:
        raise ApiError("Employee name is required.")
    salary = money_minor(data.get("salary", to_money(old["salary_paisa"])), "salary")
    join = parse_date(data.get("join_date"), "joining date") if data.get("join_date") else old["join_date"]
    status = str(data.get("status", old["status"]))
    if status not in {"active", "inactive"}:
        raise ApiError("Choose active or inactive status.")
    with atomic() as conn:
        conn.execute("UPDATE employees SET name=?,phone=?,address=?,designation=?,join_date=?,salary_paisa=?,status=?,updated_at=? WHERE id=?",
                     (name, str(data.get("phone", old["phone"] or "")).strip(), str(data.get("address", old["address"] or "")).strip(), str(data.get("designation", old["designation"] or "")).strip(), join, salary, status, now_text(), employee_id))
        audit(conn, "updated", "employees", employee_id, dict(old), {"name": name, "salary": to_money(salary), "status": status})
    return jsonify({"ok": True})


# ---------- Bank, deposits and cashbook ----------
@app.get("/api/banks")
@permission_required("cashbook:view")
def list_banks():
    rows = fetchall("SELECT * FROM bank_accounts ORDER BY bank_name")
    return jsonify([dict(r) | {"opening_balance": to_money(r["opening_balance_paisa"]), "balance": to_money(bank_balance(r["id"]))} for r in rows])


@app.post("/api/banks")
@permission_required("banks:manage")
def create_bank():
    data = request_data()
    bank_name = str(data.get("bank_name", "")).strip()
    account_name = str(data.get("account_name", "")).strip()
    if not bank_name or not account_name:
        raise ApiError("Bank name and account name are required.")
    opening = money_minor(data.get("opening_balance", 0), "opening balance")
    stamp = now_text()
    with atomic() as conn:
        cur = conn.execute("INSERT INTO bank_accounts(bank_name,account_name,account_number,opening_balance_paisa,status,created_at,updated_at) VALUES(?,?,?,?, 'active',?,?)",
                           (bank_name, account_name, str(data.get("account_number", "")).strip(), opening, stamp, stamp))
        audit(conn, "created", "banks", cur.lastrowid, None, {"bank_name": bank_name, "account_name": account_name, "opening_balance": to_money(opening)})
    return jsonify({"id": cur.lastrowid}), 201


@app.put("/api/banks/<int:bank_id>")
@permission_required("banks:manage")
def update_bank(bank_id: int):
    old = active_record("bank_accounts", bank_id)
    data = request_data()
    name = str(data.get("bank_name", old["bank_name"])).strip()
    account = str(data.get("account_name", old["account_name"])).strip()
    status = str(data.get("status", old["status"]))
    if not name or not account or status not in {"active", "inactive"}:
        raise ApiError("Enter valid bank account details.")
    with atomic() as conn:
        conn.execute("UPDATE bank_accounts SET bank_name=?,account_name=?,account_number=?,status=?,updated_at=? WHERE id=?",
                     (name, account, str(data.get("account_number", old["account_number"] or "")).strip(), status, now_text(), bank_id))
        audit(conn, "updated", "banks", bank_id, {"bank_name": old["bank_name"], "account_name": old["account_name"], "status": old["status"]}, {"bank_name": name, "account_name": account, "status": status})
    return jsonify({"ok": True})


@app.post("/api/banks/deposits")
@permission_required("banks:manage")
def bank_deposit():
    data = request_data()
    bank_id = parse_int(data.get("bank_id"), "bank")
    active_record("bank_accounts", bank_id)
    amount = money_minor(data.get("amount"), "deposit amount")
    ensure_positive(amount, "Deposit")
    day = parse_date(data.get("date"), "date")
    shift_id = parse_int(data.get("shift_id"), "shift", None)
    if shift_id: active_record("shifts", shift_id)
    slip = str(data.get("slip_no", "")).strip()
    before_cash = cash_balance()
    if amount > before_cash:
        raise ApiError(f"Deposit exceeds available cash of ৳{to_money(before_cash):,.2f}.")
    depositor = str(data.get("depositor", g.user["full_name"])).strip()
    with atomic() as conn:
        cur = conn.execute("INSERT INTO bank_transactions(date,bank_id,channel,direction,type,amount_paisa,slip_no,depositor,reference,notes,reference_type,shift_id,user_id,created_at) VALUES(?,?,'bank','in','cash_deposit',?,?,?,?,?,'cash_deposit',?,?,?)",
                           (day, bank_id, amount, slip, depositor, str(data.get("reference", "")), str(data.get("notes", "")), shift_id, g.user["id"], now_text()))
        conn.execute("INSERT INTO cash_transactions(date,shift_id,direction,type,amount_paisa,reference_type,reference_id,notes,user_id,created_at) VALUES(?,?,'out','bank_deposit',?,'bank_deposit',?,?,?,?)",
                     (day, shift_id, amount, cur.lastrowid, f"Slip {slip} · {str(data.get('notes', ''))}".strip(), g.user["id"], now_text()))
        audit(conn, "deposit_recorded", "banks", cur.lastrowid, {"cash_balance": to_money(before_cash)}, {"amount": to_money(amount), "bank_id": bank_id, "slip_no": slip}, data.get("notes"))
    return jsonify({"id": cur.lastrowid, "cash_balance": to_money(before_cash - amount), "bank_balance": to_money(bank_balance(bank_id))}), 201


@app.post("/api/banks/transactions")
@permission_required("banks:manage")
def bank_transaction():
    data = request_data()
    bank_id = parse_int(data.get("bank_id"), "bank")
    bank = active_record("bank_accounts", bank_id)
    amount = money_minor(data.get("amount"), "amount")
    ensure_positive(amount, "Amount")
    direction = str(data.get("direction", "out"))
    if direction not in {"in", "out"}:
        raise ApiError("Direction must be income or withdrawal.")
    day = parse_date(data.get("date"), "date")
    shift_id = parse_int(data.get("shift_id"), "shift", None)
    if shift_id: active_record("shifts", shift_id)
    if direction == "out" and amount > bank_balance(bank_id):
        raise ApiError("Transaction exceeds the selected account balance.")
    with atomic() as conn:
        cur = conn.execute("INSERT INTO bank_transactions(date,bank_id,channel,direction,type,amount_paisa,reference,notes,reference_type,shift_id,user_id,created_at) VALUES(?,?,'bank',?,'adjustment',?,?,?,'bank_adjustment',?,?,?)",
                           (day, bank_id, direction, amount, str(data.get("reference", "")), str(data.get("notes", "")), shift_id, g.user["id"], now_text()))
        audit(conn, "bank_transaction", "banks", cur.lastrowid, None, {"bank": bank["bank_name"], "direction": direction, "amount": to_money(amount)})
    return jsonify({"id": cur.lastrowid, "balance": to_money(bank_balance(bank_id))}), 201


@app.post("/api/cashbook/transactions")
@permission_required("cashbook:manage")
def cashbook_transaction():
    data = request_data()
    kind = str(data.get("type", "other_income"))
    if kind not in {"other_income", "withdrawal"}:
        raise ApiError("Choose other income or withdrawal.")
    amount = money_minor(data.get("amount"), "amount")
    ensure_positive(amount, "Amount")
    direction = "in" if kind == "other_income" else "out"
    day = parse_date(data.get("date"), "date")
    shift_id = parse_int(data.get("shift_id"), "shift", None)
    if shift_id: active_record("shifts", shift_id)
    if direction == "out" and amount > cash_balance():
        raise ApiError("Withdrawal exceeds available cash.")
    with atomic() as conn:
        cur = conn.execute("INSERT INTO cash_transactions(date,shift_id,direction,type,amount_paisa,notes,user_id,created_at) VALUES(?,?,?,?,?,?,?,?)",
                           (day, shift_id, direction, kind, amount, str(data.get("notes", "")), g.user["id"], now_text()))
        audit(conn, "cashbook_entry", "cashbook", cur.lastrowid, None, {"type": kind, "amount": to_money(amount)}, data.get("notes"))
    return jsonify({"id": cur.lastrowid, "cash_balance": to_money(cash_balance())}), 201


@app.get("/api/cashbook")
@permission_required("cashbook:view")
def get_cashbook():
    start, end = date_range_params()
    rows = fetchall("SELECT c.*,u.full_name AS user_name,sh.name AS shift_name FROM cash_transactions c LEFT JOIN users u ON u.id=c.user_id LEFT JOIN shifts sh ON sh.id=c.shift_id WHERE c.date BETWEEN ? AND ? ORDER BY c.date DESC,c.id DESC LIMIT 2000", (start, end))
    result = [{**dict(r), "amount": to_money(r["amount_paisa"])} for r in rows]
    totals = fetchone("SELECT COALESCE(SUM(CASE WHEN direction='in' THEN amount_paisa ELSE 0 END),0) income,COALESCE(SUM(CASE WHEN direction='out' THEN amount_paisa ELSE 0 END),0) expenses FROM cash_transactions WHERE date BETWEEN ? AND ?", (start, end))
    return jsonify({"rows": result, "totals": {"income": to_money(totals["income"]), "outflow": to_money(totals["expenses"]), "balance": to_money(cash_balance())}, "opening_cash": to_money(money_minor(setting("opening_cash", "0")))})


@app.get("/api/bank-transactions")
@permission_required("cashbook:view")
def get_bank_transactions():
    start, end = date_range_params()
    rows = fetchall("SELECT t.*,b.bank_name,b.account_name,u.full_name AS user_name,sh.name AS shift_name FROM bank_transactions t LEFT JOIN bank_accounts b ON b.id=t.bank_id LEFT JOIN users u ON u.id=t.user_id LEFT JOIN shifts sh ON sh.id=t.shift_id WHERE t.date BETWEEN ? AND ? ORDER BY t.date DESC,t.id DESC LIMIT 2000", (start, end))
    return jsonify([{**dict(r), "amount": to_money(r["amount_paisa"])} for r in rows])

# ---------- Sales, digital cash/credit memo and meter readings ----------
def build_sale_payload(data: dict, editing: sqlite3.Row | None = None) -> dict:
    items = data.get("items")
    if not isinstance(items, list) or not items:
        raise ApiError("Add at least one product to the memo.")
    clean_items = []
    total = 0
    aggregated: dict[int, float] = {}
    permissions = user_permissions(g.user)
    for item in items:
        if not isinstance(item, dict):
            raise ApiError("Each memo item must include a product, quantity and rate.")
        product_id = parse_int(item.get("product_id"), "product")
        product = active_record("products", product_id)
        if not product["active"]:
            raise ApiError(f"{product['name']} is inactive and cannot be sold.")
        qty = quantity(item.get("quantity"), "quantity")
        ensure_positive(qty, "Quantity")
        rate = money_minor(item.get("rate", to_money(product["selling_price_paisa"])), "selling rate")
        if rate != int(product["selling_price_paisa"]) and "prices:manage" not in permissions:
            raise ApiError(f"Only an authorized user can override the current price for {product['name']}.", 403, "permission_denied")
        line_total = amount_for_quantity(qty, rate)
        total += line_total
        aggregated[product_id] = aggregated.get(product_id, 0) + qty
        clean_items.append({"product": product, "product_id": product_id, "quantity": qty, "rate_paisa": rate,
                            "line_total_paisa": line_total, "unit_cost_paisa": int(product["purchase_price_paisa"])})
    customer_id = parse_int(data.get("customer_id"), "customer", None)
    if customer_id:
        customer = active_record("customers", customer_id)
        if customer["status"] != "active":
            raise ApiError("This customer is inactive.")
    else:
        customer = None
    bank_account_id = parse_int(data.get("bank_account_id"), "bank account", None)
    payment_type = str(data.get("payment_type", "cash"))
    if payment_type not in PAYMENT_METHODS:
        raise ApiError("Choose a valid payment type.")
    cash = money_minor(data.get("paid_cash", 0), "cash paid")
    bank = money_minor(data.get("paid_bank", 0), "bank paid")
    mobile = money_minor(data.get("paid_mobile", 0), "mobile paid")
    credit = money_minor(data.get("credit", 0), "credit amount")
    if cash + bank + mobile + credit == 0 and total:
        if payment_type == "credit": credit = total
        elif payment_type == "bank": bank = total
        elif payment_type == "mobile": mobile = total
        else: cash = total
    if cash + bank + mobile + credit != total:
        raise ApiError(f"Payment split must equal the memo total of ৳{to_money(total):,.2f}.")
    if credit and not customer_id:
        raise ApiError("Select a customer before recording a credit sale.")
    if (bank or mobile) and not bank_account_id:
        raise ApiError("Select the bank or mobile account receiving this payment.")
    if bank_account_id:
        bank_account = active_record("bank_accounts", bank_account_id)
        if bank_account["status"] != "active":
            raise ApiError("The selected bank account is inactive.")
    actual_type = "mixed" if sum(bool(x) for x in (cash, bank, mobile, credit)) > 1 else ("cash" if cash else "bank" if bank else "mobile" if mobile else "credit")
    sale_day = parse_date(data.get("date"), "date")
    sale_time = str(data.get("time", now_dt().strftime("%H:%M"))).strip()
    try:
        sale_time = datetime.strptime(sale_time, "%H:%M").strftime("%H:%M:%S")
    except ValueError:
        try: sale_time = datetime.strptime(sale_time, "%H:%M:%S").strftime("%H:%M:%S")
        except ValueError: raise ApiError("Time must be in HH:MM format.")
    pump_id = parse_int(data.get("pump_id"), "pump", None)
    nozzle_id = parse_int(data.get("nozzle_id"), "nozzle", None)
    shift_id = parse_int(data.get("shift_id"), "shift", None)
    operator_id = parse_int(data.get("operator_id"), "operator", None)
    if pump_id: active_record("pumps", pump_id)
    if nozzle_id:
        nozzle = active_record("nozzles", nozzle_id)
        if pump_id and nozzle["pump_id"] != pump_id:
            raise ApiError("The selected nozzle does not belong to this pump.")
        if nozzle["product_id"] not in aggregated:
            raise ApiError("Include the nozzle's assigned product in the memo.")
    if shift_id: active_record("shifts", shift_id)
    if operator_id: active_record("employees", operator_id)
    for product_id, qty in aggregated.items():
        available = float(fetchone("SELECT current_stock FROM products WHERE id=?", (product_id,))["current_stock"])
        if editing:
            old_qty = fetchone("SELECT COALESCE(SUM(quantity),0) amount FROM sale_items WHERE sale_id=? AND product_id=?", (editing["id"], product_id))["amount"]
            available += float(old_qty or 0)
        if qty > available + 0.0005:
            p = fetchone("SELECT name,unit FROM products WHERE id=?", (product_id,))
            raise ApiError(f"Insufficient {p['name']} stock. Available: {available:,.3f} {p['unit']}.", 409, "insufficient_stock")
    return {"items": clean_items, "aggregated": aggregated, "total": total, "customer_id": customer_id,
            "vehicle_no": str(data.get("vehicle_no", customer["vehicle_no"] if customer else "")).strip(),
            "payment_type": actual_type, "paid_cash": cash, "paid_bank": bank, "paid_mobile": mobile, "credit": credit,
            "bank_account_id": bank_account_id, "date": sale_day, "time": sale_time,
            "pump_id": pump_id, "nozzle_id": nozzle_id, "shift_id": shift_id, "operator_id": operator_id,
            "notes": str(data.get("notes", "")).strip()}


def post_sale_version(conn: sqlite3.Connection, sale_id: int, revision: int, payload: dict) -> None:
    source = f"sale_v{revision}"
    for item in payload["items"]:
        product = fetchone("SELECT * FROM products WHERE id=?", (item["product_id"],), conn)
        updated_stock = round(float(product["current_stock"]) - item["quantity"], 3)
        if updated_stock < -0.0005:
            raise ApiError(f"Insufficient stock for {product['name']}.", 409, "insufficient_stock")
        conn.execute("UPDATE products SET current_stock=?,updated_at=? WHERE id=?", (updated_stock, now_text(), product["id"]))
        conn.execute("INSERT INTO sale_items(sale_id,product_id,quantity,unit_price_paisa,line_total_paisa,unit_cost_paisa) VALUES(?,?,?,?,?,?)",
                     (sale_id, item["product_id"], item["quantity"], item["rate_paisa"], item["line_total_paisa"], product["purchase_price_paisa"]))
        conn.execute("INSERT INTO stock_movements(product_id,kind,quantity,business_date,unit_cost_paisa,reference_type,reference_id,notes,user_id,created_at) VALUES(?,'sale',?,?,?,?,?,?,?,?)",
                     (item["product_id"], -item["quantity"], payload["date"], product["purchase_price_paisa"], source, sale_id, f"Memo {fetchone('SELECT memo_no FROM sales WHERE id=?',(sale_id,),conn)['memo_no']}", g.user["id"], now_text()))
        if payload["nozzle_id"]:
            nozzle = fetchone("SELECT product_id,tank_id FROM nozzles WHERE id=?", (payload["nozzle_id"],), conn)
            if nozzle and nozzle["tank_id"] and nozzle["product_id"] == item["product_id"]:
                tank_qty = tank_book_quantity(nozzle["tank_id"], conn)
                if item["quantity"] > tank_qty + 0.0005:
                    tank = fetchone("SELECT tank_number FROM tanks WHERE id=?", (nozzle["tank_id"],), conn)
                    raise ApiError(f"Insufficient book quantity in tank {tank['tank_number']}.", 409, "insufficient_tank_stock")
                conn.execute("INSERT INTO tank_movements(tank_id,date,kind,quantity,reference_type,reference_id,notes,user_id,created_at) VALUES(?,?,'sale',?,'sale',?,?,?,?)",
                             (nozzle["tank_id"], payload["date"], -item["quantity"], sale_id, source, g.user["id"], now_text()))
    if payload["credit"] and payload["customer_id"]:
        conn.execute("INSERT INTO customer_transactions(customer_id,date,type,amount_paisa,source,source_id,notes,user_id,created_at) VALUES(?,?,'credit_sale',? ,?,?,?, ?,?)",
                     (payload["customer_id"], payload["date"], payload["credit"], source, sale_id, f"Credit sale memo {fetchone('SELECT memo_no FROM sales WHERE id=?',(sale_id,),conn)['memo_no']}", g.user["id"], now_text()))
    if payload["paid_cash"]:
        conn.execute("INSERT INTO cash_transactions(date,shift_id,direction,type,amount_paisa,reference_type,reference_id,notes,user_id,created_at) VALUES(?,?,'in','sale_cash',?,?,?,?,?,?)",
                     (payload["date"], payload["shift_id"], payload["paid_cash"], source, sale_id, "Cash sale", g.user["id"], now_text()))
    for channel, amount in (("bank", payload["paid_bank"]), ("mobile", payload["paid_mobile"])):
        if amount:
            conn.execute("INSERT INTO bank_transactions(date,bank_id,channel,direction,type,amount_paisa,reference,notes,reference_type,reference_id,shift_id,user_id,created_at) VALUES(?,?,?,'in','sale_collection',?,?,?,?,?,?,?,?)",
                         (payload["date"], payload["bank_account_id"], channel, amount, "Fuel station memo", f"Memo {fetchone('SELECT memo_no FROM sales WHERE id=?',(sale_id,),conn)['memo_no']}", source, sale_id, payload["shift_id"], g.user["id"], now_text()))


def reverse_sale_version(conn: sqlite3.Connection, sale: sqlite3.Row, revision: int, reason: str, cancellation: bool = False) -> list[dict]:
    source = f"sale_v{revision}"
    items = fetchall("SELECT * FROM sale_items WHERE sale_id=?", (sale["id"],), conn)
    snapshot_items = []
    for item in items:
        product = fetchone("SELECT current_stock FROM products WHERE id=?", (item["product_id"],), conn)
        new_stock = round(float(product["current_stock"]) + float(item["quantity"]), 3)
        conn.execute("UPDATE products SET current_stock=?,updated_at=? WHERE id=?", (new_stock, now_text(), item["product_id"]))
        conn.execute("INSERT INTO stock_movements(product_id,kind,quantity,business_date,unit_cost_paisa,reference_type,reference_id,notes,user_id,created_at) VALUES(?,'sale_reversal',?,?,?,?,?,?,?,?)",
                     (item["product_id"], item["quantity"], sale["date"], item["unit_cost_paisa"], source + ("_cancel" if cancellation else "_edit"), sale["id"], reason, g.user["id"], now_text()))
        if sale["nozzle_id"]:
            nozzle = fetchone("SELECT product_id,tank_id FROM nozzles WHERE id=?", (sale["nozzle_id"],), conn)
            if nozzle and nozzle["tank_id"] and nozzle["product_id"] == item["product_id"]:
                conn.execute("INSERT INTO tank_movements(tank_id,date,kind,quantity,reference_type,reference_id,notes,user_id,created_at) VALUES(?,?,'sale_reversal',?,'sale',?,?,?,?)",
                             (nozzle["tank_id"], sale["date"], item["quantity"], sale["id"], source + ("_cancel" if cancellation else "_edit"), reason, g.user["id"], now_text()))
        snapshot_items.append({"product_id": item["product_id"], "quantity": item["quantity"], "unit_price": to_money(item["unit_price_paisa"]), "line_total": to_money(item["line_total_paisa"])})
    if sale["customer_id"] and sale["credit_paisa"]:
        conn.execute("INSERT INTO customer_transactions(customer_id,date,type,amount_paisa,source,source_id,notes,user_id,created_at) VALUES(?,?,'sale_reversal',?,?,?,?,?,?)",
                     (sale["customer_id"], sale["date"], -int(sale["credit_paisa"]), source + ("_cancel" if cancellation else "_edit"), sale["id"], reason, g.user["id"], now_text()))
    cash_rows = fetchall("SELECT * FROM cash_transactions WHERE reference_type=? AND reference_id=? AND type='sale_cash' AND direction='in'", (source, sale["id"]), conn)
    for row in cash_rows:
        conn.execute("INSERT INTO cash_transactions(date,shift_id,direction,type,amount_paisa,reference_type,reference_id,notes,user_id,created_at) VALUES(?,?,'out','sale_reversal',?,?,?,?,?,?)",
                     (sale["date"], sale["shift_id"], row["amount_paisa"], source + ("_cancel" if cancellation else "_edit"), sale["id"], reason, g.user["id"], now_text()))
    bank_rows = fetchall("SELECT * FROM bank_transactions WHERE reference_type=? AND reference_id=? AND type='sale_collection' AND direction='in'", (source, sale["id"]), conn)
    for row in bank_rows:
        conn.execute("INSERT INTO bank_transactions(date,bank_id,channel,direction,type,amount_paisa,reference,notes,reference_type,reference_id,shift_id,user_id,created_at) VALUES(?,?,?,'out','sale_reversal',?,?,?,?,?,?,?,?)",
                     (sale["date"], row["bank_id"], row["channel"], row["amount_paisa"], f"Memo {sale['memo_no']}", reason, source + ("_cancel" if cancellation else "_edit"), sale["id"], sale["shift_id"], g.user["id"], now_text()))
    return snapshot_items


def sale_json(sale_id: int) -> dict:
    sale = fetchone("""SELECT s.*,c.name AS customer_name,c.phone AS customer_phone,p.pump_number,n.nozzle_number,sh.name AS shift_name,
        e.name AS operator_name,u.full_name AS created_by_name,b.bank_name,b.account_name
        FROM sales s LEFT JOIN customers c ON c.id=s.customer_id LEFT JOIN pumps p ON p.id=s.pump_id
        LEFT JOIN nozzles n ON n.id=s.nozzle_id LEFT JOIN shifts sh ON sh.id=s.shift_id LEFT JOIN employees e ON e.id=s.operator_id
        LEFT JOIN users u ON u.id=s.user_id LEFT JOIN bank_accounts b ON b.id=s.bank_account_id WHERE s.id=?""", (sale_id,))
    if not sale:
        raise ApiError("Memo not found.", 404)
    result = dict(sale)
    for key in ("total_paisa", "paid_cash_paisa", "paid_bank_paisa", "paid_mobile_paisa", "credit_paisa"):
        result[key.removesuffix("_paisa")] = to_money(sale[key])
    items = fetchall("SELECT si.*,p.name AS product_name,p.unit,p.category FROM sale_items si JOIN products p ON p.id=si.product_id WHERE si.sale_id=? ORDER BY si.id", (sale_id,))
    result["items"] = [{**dict(row), "rate": to_money(row["unit_price_paisa"]), "line_total": to_money(row["line_total_paisa"]), "unit_cost": to_money(row["unit_cost_paisa"])} for row in items]
    return result


@app.get("/api/sales")
@permission_required("sales:view")
def list_sales():
    start = parse_date(request.args.get("start"), "start date") if request.args.get("start") else "0001-01-01"
    end = parse_date(request.args.get("end"), "end date") if request.args.get("end") else today_text()
    sql = """SELECT s.id,s.memo_no,s.date,s.time,s.customer_id,c.name customer_name,s.vehicle_no,
        GROUP_CONCAT(DISTINCT p.name) AS products,GROUP_CONCAT(DISTINCT si.product_id) AS product_ids,
        COALESCE(SUM(si.quantity),0) quantity,s.total_paisa,s.payment_type,s.credit_paisa,s.paid_cash_paisa,s.paid_bank_paisa,s.paid_mobile_paisa,
        e.name operator_name,s.status,s.revision FROM sales s LEFT JOIN customers c ON c.id=s.customer_id
        LEFT JOIN sale_items si ON si.sale_id=s.id LEFT JOIN products p ON p.id=si.product_id LEFT JOIN employees e ON e.id=s.operator_id WHERE s.date BETWEEN ? AND ?"""
    params: list[Any] = [start, end]
    for key, column in (("customer_id", "s.customer_id"), ("payment_type", "s.payment_type"), ("operator_id", "s.operator_id"), ("status", "s.status")):
        if request.args.get(key):
            sql += f" AND {column}=?"; params.append(request.args[key])
    if request.args.get("vehicle"):
        sql += " AND s.vehicle_no LIKE ?"; params.append(f"%{request.args['vehicle']}%")
    if request.args.get("product_id"):
        sql += " AND EXISTS(SELECT 1 FROM sale_items sx WHERE sx.sale_id=s.id AND sx.product_id=?)"; params.append(parse_int(request.args["product_id"], "product"))
    query = str(request.args.get("q", "")).strip()
    if query:
        sql += " AND (s.memo_no LIKE ? OR c.name LIKE ? OR s.vehicle_no LIKE ? OR EXISTS(SELECT 1 FROM sale_items si2 JOIN products p2 ON p2.id=si2.product_id WHERE si2.sale_id=s.id AND p2.name LIKE ?))"
        params.extend([f"%{query}%"] * 4)
    sql += " GROUP BY s.id ORDER BY s.date DESC,s.time DESC,s.id DESC LIMIT 2000"
    rows = fetchall(sql, tuple(params))
    result = []
    for r in rows:
        item = dict(r)
        for key in ("total_paisa", "credit_paisa", "paid_cash_paisa", "paid_bank_paisa", "paid_mobile_paisa"):
            item[key.removesuffix("_paisa")] = to_money(r[key])
        result.append(item)
    return jsonify(result)


@app.get("/api/sales/<int:sale_id>")
@permission_required("sales:view")
def get_sale(sale_id: int):
    return jsonify(sale_json(sale_id))


@app.post("/api/sales")
@permission_required("sales:create")
def create_sale():
    data = request_data()
    payload = build_sale_payload(data)
    with atomic() as conn:
        number = next_number(conn, "memo_next")
        memo_no = format_memo_number(number)
        if fetchone("SELECT 1 FROM sales WHERE memo_no=?", (memo_no,), conn):
            raise ApiError(f"Memo number {memo_no} already exists. Review memo numbering in Settings.", 409, "duplicate_memo")
        stamp = now_text()
        cur = conn.execute("INSERT INTO sales(memo_no,date,time,customer_id,vehicle_no,pump_id,nozzle_id,shift_id,payment_type,total_paisa,paid_cash_paisa,paid_bank_paisa,paid_mobile_paisa,credit_paisa,bank_account_id,operator_id,user_id,notes,status,revision,created_at,updated_at) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?, 'posted',1,?,?)",
                           (memo_no, payload["date"], payload["time"], payload["customer_id"], payload["vehicle_no"], payload["pump_id"], payload["nozzle_id"], payload["shift_id"], payload["payment_type"], payload["total"], payload["paid_cash"], payload["paid_bank"], payload["paid_mobile"], payload["credit"], payload["bank_account_id"], payload["operator_id"], g.user["id"], payload["notes"], stamp, stamp))
        sale_id = cur.lastrowid
        post_sale_version(conn, sale_id, 1, payload)
        audit(conn, "created", "sales", sale_id, None, {"memo_no": memo_no, "date": payload["date"], "customer_id": payload["customer_id"], "total": to_money(payload["total"]), "payment_type": payload["payment_type"], "credit": to_money(payload["credit"]), "items": [{"product_id": i["product_id"], "quantity": i["quantity"], "rate": to_money(i["rate_paisa"])} for i in payload["items"]]})
    return jsonify(sale_json(sale_id)), 201


@app.put("/api/sales/<int:sale_id>")
@permission_required("sales:edit")
def edit_sale(sale_id: int):
    existing = fetchone("SELECT * FROM sales WHERE id=?", (sale_id,))
    if not existing:
        raise ApiError("Memo not found.", 404)
    if existing["status"] != "posted":
        raise ApiError("Cancelled memos cannot be edited.", 409)
    data = request_data()
    payload = build_sale_payload(data, existing)
    old_customer_id = existing["customer_id"]
    old_credit = int(existing["credit_paisa"])
    if old_customer_id:
        projected = balance_customer(old_customer_id) - old_credit
        if payload["customer_id"] == old_customer_id:
            projected += payload["credit"]
        if projected < 0:
            raise ApiError("This edit would create a customer advance. Record a refund/adjustment first.")
    if payload["customer_id"] and payload["customer_id"] != old_customer_id and balance_customer(payload["customer_id"]) + payload["credit"] < 0:
        raise ApiError("Customer balance cannot be negative.")
    old_snapshot = sale_json(sale_id)
    reason = str(data.get("reason", "Memo correction")).strip() or "Memo correction"
    with atomic() as conn:
        current = fetchone("SELECT * FROM sales WHERE id=?", (sale_id,), conn)
        revision = int(current["revision"])
        conn.execute("INSERT INTO sale_revisions(sale_id,revision,changed_at,user_id,snapshot,reason) VALUES(?,?,?,?,?,?)",
                     (sale_id, revision, now_text(), g.user["id"], json.dumps(old_snapshot, ensure_ascii=False, default=str), reason))
        reverse_sale_version(conn, current, revision, reason, False)
        conn.execute("DELETE FROM sale_items WHERE sale_id=?", (sale_id,))
        new_revision = revision + 1
        conn.execute("UPDATE sales SET date=?,time=?,customer_id=?,vehicle_no=?,pump_id=?,nozzle_id=?,shift_id=?,payment_type=?,total_paisa=?,paid_cash_paisa=?,paid_bank_paisa=?,paid_mobile_paisa=?,credit_paisa=?,bank_account_id=?,operator_id=?,notes=?,revision=?,updated_at=? WHERE id=?",
                     (payload["date"], payload["time"], payload["customer_id"], payload["vehicle_no"], payload["pump_id"], payload["nozzle_id"], payload["shift_id"], payload["payment_type"], payload["total"], payload["paid_cash"], payload["paid_bank"], payload["paid_mobile"], payload["credit"], payload["bank_account_id"], payload["operator_id"], payload["notes"], new_revision, now_text(), sale_id))
        post_sale_version(conn, sale_id, new_revision, payload)
        audit(conn, "edited", "sales", sale_id, old_snapshot, {"memo_no": existing["memo_no"], "total": to_money(payload["total"]), "payment_type": payload["payment_type"], "credit": to_money(payload["credit"]), "revision": new_revision}, reason)
    return jsonify(sale_json(sale_id))


@app.post("/api/sales/<int:sale_id>/cancel")
@permission_required("sales:cancel")
def cancel_sale(sale_id: int):
    data = request_data()
    reason = str(data.get("reason", "")).strip()
    if len(reason) < 5:
        raise ApiError("Please provide a cancellation reason (at least 5 characters).")
    sale = fetchone("SELECT * FROM sales WHERE id=?", (sale_id,))
    if not sale:
        raise ApiError("Memo not found.", 404)
    if sale["status"] != "posted":
        raise ApiError("This memo has already been cancelled.", 409)
    old_snapshot = sale_json(sale_id)
    with atomic() as conn:
        current = fetchone("SELECT * FROM sales WHERE id=?", (sale_id,), conn)
        reverse_sale_version(conn, current, int(current["revision"]), reason, True)
        conn.execute("UPDATE sales SET status='cancelled',updated_at=? WHERE id=?", (now_text(), sale_id))
        audit(conn, "cancelled", "sales", sale_id, old_snapshot, {"status": "cancelled"}, reason)
    return jsonify({"ok": True, "status": "cancelled"})


@app.get("/api/meter-readings")
@permission_required("inventory:view")
def list_meter_readings():
    start = parse_date(request.args.get("start"), "start date") if request.args.get("start") else "0001-01-01"
    end = parse_date(request.args.get("end"), "end date") if request.args.get("end") else today_text()
    rows = fetchall("""SELECT m.*,p.pump_number,n.nozzle_number,pr.name AS product_name,pr.unit,s.name AS shift_name,e.name AS operator_name,u.full_name AS created_by_name
        FROM meter_readings m JOIN pumps p ON p.id=m.pump_id JOIN nozzles n ON n.id=m.nozzle_id JOIN products pr ON pr.id=m.product_id
        LEFT JOIN shifts s ON s.id=m.shift_id LEFT JOIN employees e ON e.id=m.operator_id LEFT JOIN users u ON u.id=m.user_id
        WHERE m.date BETWEEN ? AND ? ORDER BY m.date DESC,m.id DESC LIMIT 2000""", (start, end))
    return jsonify([dict(r) | {"rate": to_money(r["rate_paisa"]), "total": to_money(r["total_paisa"])} for r in rows])


@app.post("/api/meter-readings")
@permission_required("meters:manage")
def create_meter_reading():
    data = request_data()
    day = parse_date(data.get("date"), "date")
    pump_id = parse_int(data.get("pump_id"), "pump")
    nozzle_id = parse_int(data.get("nozzle_id"), "nozzle")
    pump = active_record("pumps", pump_id)
    nozzle = active_record("nozzles", nozzle_id)
    if nozzle["pump_id"] != pump_id:
        raise ApiError("The selected nozzle does not belong to the selected pump.")
    opening = meter_value(data.get("opening_meter"), "opening meter")
    closing = meter_value(data.get("closing_meter"), "closing meter")
    if closing < opening:
        raise ApiError("Closing meter cannot be below opening meter.", 422, "negative_meter_reading")
    previous = fetchone("SELECT closing_meter,date FROM meter_readings WHERE nozzle_id=? ORDER BY date DESC,id DESC LIMIT 1", (nozzle_id,))
    mismatch = previous and abs(float(previous["closing_meter"]) - opening) > 0.001
    if mismatch and not data.get("confirm_mismatch"):
        raise ApiError(f"Opening meter does not match the previous closing reading ({previous['closing_meter']:,.3f} on {previous['date']}). Confirm to continue.", 409, "meter_mismatch")
    product = active_record("products", nozzle["product_id"])
    liters = round(closing - opening, 3)
    rate = int(product["selling_price_paisa"])
    total = amount_for_quantity(liters, rate)
    shift_id = parse_int(data.get("shift_id"), "shift", None)
    operator_id = parse_int(data.get("operator_id"), "operator", None)
    if shift_id: active_record("shifts", shift_id)
    if operator_id: active_record("employees", operator_id)
    with atomic() as conn:
        cur = conn.execute("INSERT INTO meter_readings(date,shift_id,pump_id,nozzle_id,product_id,opening_meter,closing_meter,sales_liters,rate_paisa,total_paisa,operator_id,user_id,notes,created_at) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
                           (day, shift_id, pump_id, nozzle_id, product["id"], opening, closing, liters, rate, total, operator_id, g.user["id"], str(data.get("notes", "")), now_text()))
        conn.execute("UPDATE nozzles SET last_meter=? WHERE id=?", (closing, nozzle_id))
        audit(conn, "meter_reading_recorded", "meters", cur.lastrowid, {"previous_closing": previous["closing_meter"] if previous else None}, {"opening": opening, "closing": closing, "sales_liters": liters, "amount": to_money(total)}, "opening mismatch confirmed" if mismatch else None)
    return jsonify({"id": cur.lastrowid, "sales_liters": liters, "rate": to_money(rate), "total": to_money(total), "meter_mismatch_confirmed": bool(mismatch)}), 201

# ---------- Fuel purchases and expenses ----------
@app.get("/api/purchases")
@permission_required("purchases:view")
def list_purchases():
    start = parse_date(request.args.get("start"), "start date") if request.args.get("start") else "0001-01-01"
    end = parse_date(request.args.get("end"), "end date") if request.args.get("end") else today_text()
    query = str(request.args.get("q", "")).strip()
    sql = """SELECT pu.*,s.name AS supplier_name,p.name AS product_name,p.unit,t.tank_number,u.full_name AS created_by_name,b.bank_name
        FROM purchases pu JOIN suppliers s ON s.id=pu.supplier_id JOIN products p ON p.id=pu.product_id
        LEFT JOIN tanks t ON t.id=pu.tank_id LEFT JOIN users u ON u.id=pu.user_id LEFT JOIN bank_accounts b ON b.id=pu.bank_account_id
        WHERE pu.date BETWEEN ? AND ?"""
    params: list[Any] = [start, end]
    if query:
        sql += " AND (pu.invoice_no LIKE ? OR pu.challan_no LIKE ? OR s.name LIKE ? OR p.name LIKE ?)"
        params.extend([f"%{query}%"] * 4)
    sql += " ORDER BY pu.date DESC,pu.id DESC LIMIT 2000"
    money_fields = ("purchase_rate_paisa", "total_paisa", "paid_amount_paisa", "due_amount_paisa", "transport_cost_paisa", "other_cost_paisa")
    return jsonify([{**dict(r), **{key.removesuffix("_paisa"): to_money(r[key]) for key in money_fields}} for r in fetchall(sql, tuple(params))])


@app.post("/api/purchases")
@permission_required("purchases:create")
def create_purchase():
    data = request_data()
    supplier_id = parse_int(data.get("supplier_id"), "supplier")
    supplier = active_record("suppliers", supplier_id)
    if supplier["status"] != "active":
        raise ApiError("The selected supplier is inactive.")
    product_id = parse_int(data.get("product_id"), "product")
    product = active_record("products", product_id)
    qty = quantity(data.get("quantity"), "quantity")
    ensure_positive(qty, "Purchase quantity")
    rate = money_minor(data.get("purchase_rate", to_money(product["purchase_price_paisa"])), "purchase rate")
    transport = money_minor(data.get("transport_cost", 0), "transport cost")
    other = money_minor(data.get("other_cost", 0), "other cost")
    total = amount_for_quantity(qty, rate) + transport + other
    paid = money_minor(data.get("paid_amount", 0), "paid amount")
    if paid > total:
        raise ApiError("Paid amount cannot be greater than the purchase total.")
    method = str(data.get("payment_method", "cash"))
    if method not in {"cash", "bank", "mobile", "credit"}:
        raise ApiError("Choose a valid purchase payment method.")
    if method == "credit" and paid:
        raise ApiError("Use cash, bank or mobile for an immediate payment; credit means the full amount is unpaid.")
    if method in {"bank", "mobile"} and paid:
        account_id = parse_int(data.get("bank_account_id"), "bank account", None)
        if not account_id:
            raise ApiError("Select the account used to pay the supplier.")
        bank_account = active_record("bank_accounts", account_id)
        if bank_account["status"] != "active":
            raise ApiError("The selected account is inactive.")
    else:
        account_id = parse_int(data.get("bank_account_id"), "bank account", None)
        if account_id: active_record("bank_accounts", account_id)
    if method == "cash" and paid > cash_balance():
        raise ApiError("Purchase payment exceeds available cash. Record an opening balance or cash receipt first.")
    day = parse_date(data.get("date"), "date")
    shift_id = parse_int(data.get("shift_id"), "shift", None)
    if shift_id: active_record("shifts", shift_id)
    tank_id = parse_int(data.get("tank_id"), "tank", None)
    if tank_id:
        tank = active_record("tanks", tank_id)
        if tank["product_id"] != product_id:
            raise ApiError("The selected tank is assigned to a different product.")
        tank_balance = tank_book_quantity(tank_id)
        if tank_balance + qty > float(tank["capacity"]) + 0.0005:
            raise ApiError(f"Purchase would exceed {tank['tank_number']} capacity. Book tank quantity is {tank_balance:,.3f} and capacity is {float(tank['capacity']):,.3f}.")
    with atomic() as conn:
        seq = next_number(conn, "purchase_next")
        invoice = f"{setting('purchase_prefix', 'PUR-')}{seq}"
        if fetchone("SELECT 1 FROM purchases WHERE invoice_no=?", (invoice,), conn):
            raise ApiError("Purchase invoice number already exists. Review Settings.", 409, "duplicate_invoice")
        stamp = now_text()
        cur = conn.execute("INSERT INTO purchases(invoice_no,date,supplier_id,product_id,quantity,purchase_rate_paisa,total_paisa,paid_amount_paisa,due_amount_paisa,transport_cost_paisa,other_cost_paisa,payment_method,bank_account_id,tank_id,challan_no,notes,status,user_id,created_at) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?, 'posted',?,?)",
                           (invoice, day, supplier_id, product_id, qty, rate, total, paid, total - paid, transport, other, method, account_id, tank_id, str(data.get("challan_no", "")).strip(), str(data.get("notes", "")), g.user["id"], stamp))
        purchase_id = cur.lastrowid
        # Moving weighted-average inventory costing includes freight and other landed costs.
        current_product = fetchone("SELECT * FROM products WHERE id=?", (product_id,), conn)
        landed_unit_cost = int((Decimal(total) / Decimal(str(qty))).quantize(Decimal("1"), rounding=ROUND_HALF_UP))
        current_qty = Decimal(str(current_product["current_stock"]))
        if current_qty + Decimal(str(qty)) > 0:
            weighted = ((current_qty * Decimal(int(current_product["purchase_price_paisa"]))) + (Decimal(str(qty)) * Decimal(landed_unit_cost))) / (current_qty + Decimal(str(qty)))
            new_cost = int(weighted.quantize(Decimal("1"), rounding=ROUND_HALF_UP))
        else:
            new_cost = landed_unit_cost
        new_stock = round(float(current_product["current_stock"]) + qty, 3)
        conn.execute("UPDATE products SET current_stock=?,purchase_price_paisa=?,updated_at=? WHERE id=?", (new_stock, new_cost, stamp, product_id))
        conn.execute("INSERT INTO stock_movements(product_id,kind,quantity,business_date,unit_cost_paisa,reference_type,reference_id,notes,user_id,created_at) VALUES(?,'purchase',?,?,?,?,?,?,?,?)",
                     (product_id, qty, day, landed_unit_cost, "purchase", purchase_id, f"Purchase {invoice}", g.user["id"], stamp))
        if tank_id:
            conn.execute("INSERT INTO tank_movements(tank_id,date,kind,quantity,reference_type,reference_id,notes,user_id,created_at) VALUES(?,?,'purchase',?,'purchase',?,?,?,?)",
                         (tank_id, day, qty, purchase_id, f"Purchase {invoice}", g.user["id"], stamp))
        conn.execute("INSERT INTO supplier_transactions(supplier_id,date,type,amount_paisa,source,source_id,notes,user_id,created_at) VALUES(?,?,'purchase',?,'purchase',?,?,?,?)",
                     (supplier_id, day, total, purchase_id, invoice, g.user["id"], stamp))
        if paid:
            conn.execute("INSERT INTO supplier_transactions(supplier_id,date,type,amount_paisa,source,source_id,notes,user_id,created_at) VALUES(?,?,'payment',?,'purchase_payment',?,?,?,?)",
                         (supplier_id, day, -paid, purchase_id, invoice, g.user["id"], stamp))
            if method == "cash":
                conn.execute("INSERT INTO cash_transactions(date,shift_id,direction,type,amount_paisa,reference_type,reference_id,notes,user_id,created_at) VALUES(?,?,'out','purchase_payment',?,'purchase',?,?,?,?)",
                             (day, shift_id, paid, purchase_id, f"{invoice} {str(data.get('notes', ''))}".strip(), g.user["id"], stamp))
            elif method in {"bank", "mobile"}:
                conn.execute("INSERT INTO bank_transactions(date,bank_id,channel,direction,type,amount_paisa,reference,notes,reference_type,reference_id,shift_id,user_id,created_at) VALUES(?,?,?,'out','purchase_payment',?,?,?,?,?,?,?,?)",
                             (day, account_id, method, paid, invoice, str(data.get("notes", "")), "purchase", purchase_id, shift_id, g.user["id"], stamp))
        audit(conn, "created", "purchases", purchase_id, None, {"invoice_no": invoice, "supplier": supplier["name"], "product": product["name"], "quantity": qty, "rate": to_money(rate), "total": to_money(total), "paid": to_money(paid), "due": to_money(total - paid)})
    return jsonify({"id": purchase_id, "invoice_no": invoice, "total": to_money(total), "due": to_money(total - paid), "current_stock": new_stock, "average_cost": to_money(new_cost)}), 201


@app.get("/api/expenses")
@permission_required("cashbook:view")
def list_expenses():
    start = parse_date(request.args.get("start"), "start date") if request.args.get("start") else "0001-01-01"
    end = parse_date(request.args.get("end"), "end date") if request.args.get("end") else today_text()
    sql = "SELECT x.*,u.full_name AS created_by_name,a.full_name AS approved_by_name,b.bank_name,sh.name AS shift_name FROM expenses x LEFT JOIN users u ON u.id=x.created_by LEFT JOIN users a ON a.id=x.approved_by LEFT JOIN bank_accounts b ON b.id=x.bank_account_id LEFT JOIN shifts sh ON sh.id=x.shift_id WHERE x.date BETWEEN ? AND ?"
    params: list[Any] = [start, end]
    if request.args.get("status"):
        sql += " AND x.status=?"; params.append(request.args["status"])
    sql += " ORDER BY x.date DESC,x.id DESC LIMIT 2000"
    return jsonify([{**dict(r), "amount": to_money(r["amount_paisa"])} for r in fetchall(sql, tuple(params))])


@app.post("/api/expenses")
@permission_required("expenses:manage")
def create_expense():
    data = request_data()
    amount = money_minor(data.get("amount"), "expense amount")
    ensure_positive(amount, "Expense")
    category = str(data.get("category", "Other"))
    if category not in EXPENSE_CATEGORIES:
        raise ApiError("Choose a valid expense category.")
    method = str(data.get("payment_method", "cash"))
    if method not in {"cash", "bank", "mobile", "credit"}:
        raise ApiError("Choose cash, bank, mobile or credit payment.")
    account_id = parse_int(data.get("bank_account_id"), "bank account", None)
    if method in {"bank", "mobile"} and not account_id:
        raise ApiError("Select the account used to pay this expense.")
    if account_id: active_record("bank_accounts", account_id)
    if method == "cash" and amount > cash_balance():
        raise ApiError("Expense exceeds available cash.")
    day = parse_date(data.get("date"), "date")
    shift_id = parse_int(data.get("shift_id"), "shift", None)
    if shift_id: active_record("shifts", shift_id)
    # Managers submit expenses for approval; owners and accountants can post immediately.
    status = "approved" if g.user["role"] in {"owner", "accountant"} else "pending"
    with atomic() as conn:
        cur = conn.execute("INSERT INTO expenses(date,category,description,amount_paisa,payment_method,bank_account_id,paid_to,voucher_no,shift_id,status,created_by,approved_by,created_at) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)",
                           (day, category, str(data.get("description", "")).strip(), amount, method, account_id, str(data.get("paid_to", "")).strip(), str(data.get("voucher_no", "")).strip(), shift_id, status, g.user["id"], g.user["id"] if status == "approved" else None, now_text()))
        if status == "approved":
            post_expense_payment(conn, cur.lastrowid, day, method, amount, account_id, category, shift_id)
        audit(conn, "created", "expenses", cur.lastrowid, None, {"category": category, "amount": to_money(amount), "payment_method": method, "status": status})
    return jsonify({"id": cur.lastrowid, "status": status}), 201


def post_expense_payment(conn: sqlite3.Connection, expense_id: int, day: str, method: str, amount: int, bank_id: int | None, category: str, shift_id: int | None):
    if method == "cash":
        conn.execute("INSERT INTO cash_transactions(date,shift_id,direction,type,amount_paisa,reference_type,reference_id,notes,user_id,created_at) VALUES(?,?,'out','expense',?,'expense',?,?,?,?)",
                     (day, shift_id, amount, expense_id, category, g.user["id"], now_text()))
    elif method in {"bank", "mobile"}:
        conn.execute("INSERT INTO bank_transactions(date,bank_id,channel,direction,type,amount_paisa,reference,notes,reference_type,reference_id,shift_id,user_id,created_at) VALUES(?,?,?,'out','expense',?,?,?,?,?,?,?,?)",
                     (day, bank_id, method, amount, category, f"Expense {expense_id}", "expense", expense_id, shift_id, g.user["id"], now_text()))


@app.post("/api/expenses/<int:expense_id>/approve")
@permission_required("expenses:manage")
def approve_expense(expense_id: int):
    expense = fetchone("SELECT * FROM expenses WHERE id=?", (expense_id,))
    if not expense:
        raise ApiError("Expense not found.", 404)
    if expense["status"] != "pending":
        raise ApiError("Only pending expenses can be approved.", 409)
    if expense["payment_method"] == "cash" and expense["amount_paisa"] > cash_balance():
        raise ApiError("Expense exceeds available cash; record cash before approving.")
    with atomic() as conn:
        conn.execute("UPDATE expenses SET status='approved',approved_by=? WHERE id=?", (g.user["id"], expense_id))
        post_expense_payment(conn, expense_id, expense["date"], expense["payment_method"], expense["amount_paisa"], expense["bank_account_id"], expense["category"], expense["shift_id"])
        audit(conn, "approved", "expenses", expense_id, {"status": "pending"}, {"status": "approved", "amount": to_money(expense["amount_paisa"])})
    return jsonify({"ok": True})


@app.post("/api/expenses/<int:expense_id>/reject")
@permission_required("expenses:manage")
def reject_expense(expense_id: int):
    data = request_data()
    reason = str(data.get("reason", "")).strip()
    if not reason:
        raise ApiError("Provide a reason for rejecting the expense.")
    expense = fetchone("SELECT * FROM expenses WHERE id=?", (expense_id,))
    if not expense:
        raise ApiError("Expense not found.", 404)
    if expense["status"] != "pending":
        raise ApiError("Only pending expenses can be rejected.", 409)
    with atomic() as conn:
        conn.execute("UPDATE expenses SET status='rejected',approved_by=? WHERE id=?", (g.user["id"], expense_id))
        audit(conn, "rejected", "expenses", expense_id, {"status": "pending"}, {"status": "rejected"}, reason)
    return jsonify({"ok": True})

@app.post("/api/expenses/<int:expense_id>/reverse")
@permission_required("expenses:manage")
def reverse_expense(expense_id: int):
    data=request_data();reason=str(data.get("reason","")).strip()
    if len(reason)<5:raise ApiError("Provide a reason for reversing this expense (at least 5 characters).")
    expense=fetchone("SELECT * FROM expenses WHERE id=?",(expense_id,))
    if not expense:raise ApiError("Expense not found.",404)
    if expense["status"]!="approved":raise ApiError("Only an approved expense can be reversed.",409)
    with atomic() as conn:
        if expense["payment_method"]=="cash":
            existing=fetchone("SELECT id FROM cash_transactions WHERE reference_type='expense' AND reference_id=? AND type='expense' AND direction='out'",(expense_id,),conn)
            if existing:
                conn.execute("INSERT INTO cash_transactions(date,shift_id,direction,type,amount_paisa,reference_type,reference_id,notes,user_id,created_at) VALUES(?,?,'in','expense_reversal',?,'expense_reverse',?,?,?,?)",(expense["date"],expense["shift_id"],expense["amount_paisa"],expense_id,reason,g.user["id"],now_text()))
        elif expense["payment_method"] in {"bank","mobile"}:
            existing=fetchone("SELECT id FROM bank_transactions WHERE reference_type='expense' AND reference_id=? AND type='expense' AND direction='out'",(expense_id,),conn)
            if existing:
                conn.execute("INSERT INTO bank_transactions(date,bank_id,channel,direction,type,amount_paisa,reference,notes,reference_type,reference_id,shift_id,user_id,created_at) VALUES(?,?,?,'in','expense_reversal',?,?,?,?,?,?,?,?)",(expense["date"],expense["bank_account_id"],expense["payment_method"],expense["amount_paisa"],f"Expense {expense_id}",reason,"expense_reverse",expense_id,expense["shift_id"],g.user["id"],now_text()))
        conn.execute("UPDATE expenses SET status='reversed' WHERE id=?",(expense_id,))
        audit(conn,"expense_reversed","expenses",expense_id,{"status":"approved","amount":to_money(expense["amount_paisa"])},{"status":"reversed"},reason)
    return jsonify({"ok":True,"status":"reversed"})


# ---------- Dashboard, stock position, reports and trends ----------
def stock_position(start: str, end: str) -> list[dict]:
    products = fetchall("SELECT * FROM products ORDER BY category,name")
    output = []
    for product in products:
        opening = fetchone("SELECT COALESCE(SUM(quantity),0) value FROM stock_movements WHERE product_id=? AND business_date<?", (product["id"], start))["value"]
        movement = fetchall("SELECT kind,COALESCE(SUM(quantity),0) value FROM stock_movements WHERE product_id=? AND business_date BETWEEN ? AND ? GROUP BY kind", (product["id"], start, end))
        totals = {r["kind"]: float(r["value"] or 0) for r in movement}
        purchased = totals.get("purchase", 0.0)
        # Reversals are signed inventory movements, so net sold is sales less reversed units.
        sold = -totals.get("sale", 0.0) - totals.get("sale_reversal", 0.0)
        adjusted = totals.get("adjustment", 0.0)
        closing = round(float(opening or 0) + purchased - sold + adjusted, 3)
        latest = fetchone("""SELECT SUM(d.physical_stock) physical_stock,MAX(d.date) date,MIN(d.verified) verified,COUNT(*) tank_count FROM tank_dips d JOIN tanks t ON t.id=d.tank_id
            WHERE t.product_id=? AND d.id=(SELECT d2.id FROM tank_dips d2 WHERE d2.tank_id=d.tank_id ORDER BY d2.date DESC,d2.id DESC LIMIT 1)""", (product["id"],))
        physical = float(latest["physical_stock"]) if latest and latest["physical_stock"] is not None else None
        output.append({"product_id": product["id"], "product": product["name"], "category": product["category"], "unit": product["unit"],
                       "opening_stock": round(float(opening or 0), 3), "purchased": round(purchased, 3), "sold": round(sold, 3),
                       "adjusted": round(adjusted, 3), "closing_book_stock": closing, "current_stock": round(float(product["current_stock"]), 3),
                       "physical_stock": physical, "difference": round(physical - float(product["current_stock"]), 3) if physical is not None else None,
                       "minimum_stock": float(product["min_stock"]), "low_stock": bool(float(product["current_stock"]) <= float(product["min_stock"]) and (product["category"] == "fuel" or float(product["min_stock"]) > 0)),
                       "active": bool(product["active"]), "latest_dip_date": latest["date"] if latest else None,
                       "physical_verified": bool(latest["verified"]) if latest else None})
    return output


def sales_revenue_and_cogs(start: str, end: str, group: str = "") -> tuple[int, int]:
    row = fetchone("""SELECT COALESCE(SUM(si.line_total_paisa),0) revenue,
        COALESCE(SUM(CAST(ROUND(si.quantity * si.unit_cost_paisa,0) AS INTEGER)),0) cogs
        FROM sale_items si JOIN sales s ON s.id=si.sale_id WHERE s.status='posted' AND s.date BETWEEN ? AND ?""", (start, end))
    return int(row["revenue"] or 0), int(row["cogs"] or 0)


def dashboard_data(start: str, end: str) -> dict:
    sales_total_row = fetchone("SELECT COALESCE(SUM(total_paisa),0) value FROM sales WHERE status='posted' AND date BETWEEN ? AND ?", (start, end))
    sales_total = int(sales_total_row["value"] or 0)
    fuel_sales = fetchone("SELECT COALESCE(SUM(si.line_total_paisa),0) value FROM sale_items si JOIN sales s ON s.id=si.sale_id JOIN products p ON p.id=si.product_id WHERE s.status='posted' AND s.date BETWEEN ? AND ? AND p.category='fuel'", (start, end))["value"]
    lubricant_sales = fetchone("SELECT COALESCE(SUM(si.line_total_paisa),0) value FROM sale_items si JOIN sales s ON s.id=si.sale_id JOIN products p ON p.id=si.product_id WHERE s.status='posted' AND s.date BETWEEN ? AND ? AND p.category='lubricant'", (start, end))["value"]
    payment = fetchone("SELECT COALESCE(SUM(paid_cash_paisa),0) cash,COALESCE(SUM(credit_paisa),0) credit FROM sales WHERE status='posted' AND date BETWEEN ? AND ?", (start, end))
    collection = fetchone("SELECT COALESCE(SUM(amount_paisa),0) value FROM cash_transactions WHERE date BETWEEN ? AND ? AND direction='in'", (start, end))["value"]
    expense = fetchone("SELECT COALESCE(SUM(amount_paisa),0) value FROM expenses WHERE status='approved' AND date BETWEEN ? AND ?", (start, end))["value"]
    deposit = fetchone("SELECT COALESCE(SUM(amount_paisa),0) value FROM bank_transactions WHERE date BETWEEN ? AND ? AND type='cash_deposit' AND direction='in'", (start, end))["value"]
    revenue, cogs = sales_revenue_and_cogs(start, end)
    due_row = fetchone("SELECT COALESCE(SUM(c.opening_due_paisa + COALESCE(t.total,0)),0) value FROM customers c LEFT JOIN (SELECT customer_id,SUM(amount_paisa) total FROM customer_transactions GROUP BY customer_id) t ON t.customer_id=c.id WHERE c.status='active'")
    stock = fetchone("SELECT COALESCE(SUM(current_stock),0) value FROM products WHERE active=1 AND category='fuel'")["value"]
    daily_sales = fetchall("SELECT date,COALESCE(SUM(total_paisa),0) value FROM sales WHERE status='posted' AND date BETWEEN ? AND ? GROUP BY date ORDER BY date", (start, end))
    product_rows = fetchall("SELECT p.name,COALESCE(SUM(si.line_total_paisa),0) value FROM sale_items si JOIN sales s ON s.id=si.sale_id JOIN products p ON p.id=si.product_id WHERE s.status='posted' AND s.date BETWEEN ? AND ? GROUP BY p.id ORDER BY value DESC LIMIT 8", (start, end))
    # The trend series are based on actual posted entries; missing dates are zero-filled, not sample data.
    end_dt = datetime.strptime(end, "%Y-%m-%d").date()
    daily_map = {r["date"]: int(r["value"] or 0) for r in daily_sales}
    daily_labels = [(end_dt - timedelta(days=i)).strftime("%Y-%m-%d") for i in range(13, -1, -1)]
    daily_chart = [{"label": d, "value": to_money(daily_map.get(d, 0))} for d in daily_labels]
    weeks = [(end_dt - timedelta(days=i)).strftime("%Y-%m-%d") for i in range(6, -1, -1)]
    week_map = {r["date"]: int(r["value"] or 0) for r in fetchall("SELECT date,COALESCE(SUM(total_paisa),0) value FROM sales WHERE status='posted' AND date BETWEEN ? AND ? GROUP BY date", ((end_dt - timedelta(days=6)).strftime("%Y-%m-%d"), end))}
    weekly_chart = [{"label": d, "value": to_money(week_map.get(d, 0))} for d in weeks]
    month_labels = []
    cursor_month = end_dt.replace(day=1)
    for _ in range(11, -1, -1):
        year = cursor_month.year
        month = cursor_month.month - _
        while month <= 0:
            month += 12; year -= 1
        month_labels.append(f"{year:04d}-{month:02d}")
    monthly_rows = fetchall("SELECT substr(date,1,7) month,COALESCE(SUM(total_paisa),0) value FROM sales WHERE status='posted' AND substr(date,1,7) BETWEEN ? AND ? GROUP BY substr(date,1,7)", (month_labels[0], month_labels[-1]))
    monthly_map = {r["month"]: int(r["value"] or 0) for r in monthly_rows}
    monthly_chart = [{"label": m, "value": to_money(monthly_map.get(m, 0))} for m in month_labels]
    expense_rows = fetchall("SELECT date,COALESCE(SUM(amount_paisa),0) value FROM expenses WHERE status='approved' AND date BETWEEN ? AND ? GROUP BY date", ((end_dt - timedelta(days=13)).strftime("%Y-%m-%d"), end))
    expense_map = {r["date"]: int(r["value"] or 0) for r in expense_rows}
    expense_chart = [{"label": d, "value": to_money(expense_map.get(d, 0))} for d in daily_labels]
    profit_rows = fetchall("""SELECT s.date,COALESCE(SUM(si.line_total_paisa-CAST(ROUND(si.quantity*si.unit_cost_paisa,0) AS INTEGER)),0) value
        FROM sales s JOIN sale_items si ON si.sale_id=s.id WHERE s.status='posted' AND s.date BETWEEN ? AND ? GROUP BY s.date""", ((end_dt - timedelta(days=13)).strftime("%Y-%m-%d"), end))
    profit_map = {r["date"]: int(r["value"] or 0) for r in profit_rows}
    profit_chart = [{"label": d, "value": to_money(profit_map.get(d, 0) - expense_map.get(d, 0))} for d in daily_labels]
    return {
        "period": {"start": start, "end": end},
        "cards": {"sales": to_money(sales_total), "fuel_sales": to_money(fuel_sales), "lubricant_sales": to_money(lubricant_sales),
                  "cash_collection": to_money(collection), "credit_sales": to_money(payment["credit"]), "customer_due": to_money(due_row["value"]),
                  "expenses": to_money(expense), "bank_deposit": to_money(deposit), "fuel_stock": round(float(stock or 0), 3),
                  "gross_profit": to_money(revenue - cogs), "net_profit": to_money(revenue - cogs - int(expense or 0)),
                  "sales_count": int(fetchone("SELECT COUNT(*) value FROM sales WHERE status='posted' AND date BETWEEN ? AND ?", (start, end))["value"])},
        "charts": {"daily_sales": daily_chart, "product_sales": [{"label": r["name"], "value": to_money(r["value"])} for r in product_rows],
                   "weekly_sales": weekly_chart, "monthly_sales": monthly_chart, "expense_trend": expense_chart, "profit_trend": profit_chart},
        "cash_balance": to_money(cash_balance()), "database": "live",
    }


@app.get("/api/dashboard")
@permission_required("dashboard:view")
def dashboard():
    start, end = date_range_params()
    return jsonify(dashboard_data(start, end))


REPORT_TITLES = {
    "daily_sales": "Daily Sales Report", "daily_closing": "Daily Closing Report", "product_sales": "Product-wise Sales Report",
    "pump_sales": "Pump-wise Sales Report", "nozzle_sales": "Nozzle-wise Sales Report", "shift_sales": "Shift-wise Sales Report",
    "meter_readings": "Meter Reading Report", "fuel_stock": "Fuel Stock Report", "tank_dips": "Tank Dip Report",
    "purchases": "Purchase Report", "supplier_due": "Supplier Due Report", "customer_due": "Customer Due Report",
    "customer_statement": "Customer Statement", "cashbook": "Cashbook Report", "bank_deposits": "Bank Deposit Report",
    "expenses": "Expense Report", "profit_loss": "Profit & Loss", "employee_activity": "Employee Activity Report",
    "audit_logs": "Audit Log Report", "price_changes": "Price Change Report",
}


def col(key: str, label: str) -> dict:
    return {"key": key, "label": label}


def make_report(report_type: str, start: str, end: str, args: dict | None = None) -> dict:
    args = args or {}
    rows: list[dict] = []
    columns: list[dict] = []
    summary: dict[str, Any] = {}
    if report_type == "daily_sales":
        columns = [col("memo_no", "Memo No"), col("date", "Date"), col("time", "Time"), col("customer_name", "Customer"), col("vehicle_no", "Vehicle"), col("product_name", "Product"), col("quantity", "Quantity"), col("unit", "Unit"), col("rate", "Rate (৳)"), col("amount", "Amount (৳)"), col("payment_type", "Payment"), col("operator_name", "Operator"), col("status", "Status")]
        raw = fetchall("""SELECT s.memo_no,s.date,s.time,c.name customer_name,s.vehicle_no,p.name product_name,si.quantity,p.unit,
            si.unit_price_paisa,si.line_total_paisa,s.payment_type,e.name operator_name,s.status
            FROM sales s JOIN sale_items si ON si.sale_id=s.id JOIN products p ON p.id=si.product_id
            LEFT JOIN customers c ON c.id=s.customer_id LEFT JOIN employees e ON e.id=s.operator_id
            WHERE s.date BETWEEN ? AND ? ORDER BY s.date DESC,s.time DESC,s.id DESC""", (start, end))
        rows = [{**dict(r), "rate": to_money(r["unit_price_paisa"]), "amount": to_money(r["line_total_paisa"])} for r in raw]
        summary["total_sales"] = to_money(sum(int(r["line_total_paisa"]) for r in raw if r["status"] == "posted"))
    elif report_type == "daily_closing":
        columns = [col("label", "Description"), col("amount", "Amount (৳)")]
        sales = fetchone("SELECT COALESCE(SUM(total_paisa),0) sales,COALESCE(SUM(paid_cash_paisa),0) cash,COALESCE(SUM(credit_paisa),0) credit FROM sales WHERE date BETWEEN ? AND ? AND status='posted'", (start, end))
        ex = fetchone("SELECT COALESCE(SUM(amount_paisa),0) amount FROM expenses WHERE date BETWEEN ? AND ? AND status='approved'", (start, end))
        deposits = fetchone("SELECT COALESCE(SUM(amount_paisa),0) amount FROM bank_transactions WHERE date BETWEEN ? AND ? AND type='cash_deposit' AND direction='in'", (start, end))
        cash_in = fetchone("SELECT COALESCE(SUM(amount_paisa),0) amount FROM cash_transactions WHERE date BETWEEN ? AND ? AND direction='in'", (start, end))
        cash_out = fetchone("SELECT COALESCE(SUM(amount_paisa),0) amount FROM cash_transactions WHERE date BETWEEN ? AND ? AND direction='out'", (start, end))
        rows = [{"label": "Total posted sales", "amount": to_money(sales["sales"])}, {"label": "Cash sales", "amount": to_money(sales["cash"])}, {"label": "Credit sales", "amount": to_money(sales["credit"])},
                {"label": "Cash received", "amount": to_money(cash_in["amount"])}, {"label": "Cash outflow", "amount": to_money(cash_out["amount"])}, {"label": "Approved expenses", "amount": to_money(ex["amount"])}, {"label": "Bank deposits", "amount": to_money(deposits["amount"])}, {"label": "Cash on hand (live)", "amount": to_money(cash_balance())}]
    elif report_type == "product_sales":
        columns = [col("product_name", "Product"), col("category", "Category"), col("unit", "Unit"), col("quantity", "Quantity"), col("amount", "Sales (৳)"), col("cogs", "Cost of Goods (৳)"), col("gross_profit", "Gross Profit (৳)")]
        raw = fetchall("""SELECT p.name product_name,p.category,p.unit,COALESCE(SUM(si.quantity),0) quantity,
            COALESCE(SUM(si.line_total_paisa),0) amount,COALESCE(SUM(CAST(ROUND(si.quantity*si.unit_cost_paisa,0) AS INTEGER)),0) cogs
            FROM sale_items si JOIN sales s ON s.id=si.sale_id JOIN products p ON p.id=si.product_id
            WHERE s.status='posted' AND s.date BETWEEN ? AND ? GROUP BY p.id ORDER BY amount DESC""", (start, end))
        rows = [{**dict(r), "amount": to_money(r["amount"]), "cogs": to_money(r["cogs"]), "gross_profit": to_money(r["amount"] - r["cogs"])} for r in raw]
    elif report_type in {"pump_sales", "nozzle_sales", "shift_sales"}:
        if report_type == "pump_sales":
            columns = [col("pump_number", "Pump"), col("quantity", "Quantity"), col("sales", "Sales (৳)"), col("memos", "Memos")]
            sql = "SELECT p.pump_number,COALESCE(SUM(si.quantity),0) quantity,COALESCE(SUM(si.line_total_paisa),0) sales,COUNT(DISTINCT s.id) memos FROM sales s JOIN sale_items si ON si.sale_id=s.id LEFT JOIN pumps p ON p.id=s.pump_id WHERE s.status='posted' AND s.date BETWEEN ? AND ? GROUP BY p.id ORDER BY p.pump_number"
        elif report_type == "nozzle_sales":
            columns = [col("pump_number", "Pump"), col("nozzle_number", "Nozzle"), col("product_name", "Product"), col("quantity", "Quantity"), col("sales", "Sales (৳)"), col("memos", "Memos")]
            sql = "SELECT p.pump_number,n.nozzle_number,pr.name product_name,COALESCE(SUM(si.quantity),0) quantity,COALESCE(SUM(si.line_total_paisa),0) sales,COUNT(DISTINCT s.id) memos FROM sales s JOIN sale_items si ON si.sale_id=s.id LEFT JOIN pumps p ON p.id=s.pump_id LEFT JOIN nozzles n ON n.id=s.nozzle_id JOIN products pr ON pr.id=si.product_id WHERE s.status='posted' AND s.date BETWEEN ? AND ? GROUP BY n.id ORDER BY p.pump_number,n.nozzle_number"
        else:
            columns = [col("shift_name", "Shift"), col("quantity", "Quantity"), col("sales", "Sales (৳)"), col("cash", "Cash (৳)"), col("credit", "Credit (৳)"), col("memos", "Memos")]
            sql = "WITH x AS (SELECT id,shift_id,paid_cash_paisa,credit_paisa FROM sales WHERE status='posted' AND date BETWEEN ? AND ?), y AS (SELECT sale_id,SUM(quantity) quantity,SUM(line_total_paisa) sales FROM sale_items GROUP BY sale_id) SELECT sh.name shift_name,COALESCE(SUM(y.quantity),0) quantity,COALESCE(SUM(y.sales),0) sales,COALESCE(SUM(x.paid_cash_paisa),0) cash,COALESCE(SUM(x.credit_paisa),0) credit,COUNT(x.id) memos FROM x LEFT JOIN y ON y.sale_id=x.id LEFT JOIN shifts sh ON sh.id=x.shift_id GROUP BY sh.id ORDER BY sh.name"
        raw = fetchall(sql, (start, end))
        rows = []
        for r in raw:
            item = dict(r)
            for key in ("sales", "cash", "credit"):
                if key in item and item[key] is not None: item[key] = to_money(item[key])
            rows.append(item)
    elif report_type == "meter_readings":
        columns = [col("date", "Date"), col("shift_name", "Shift"), col("pump_number", "Pump"), col("nozzle_number", "Nozzle"), col("product_name", "Product"), col("opening_meter", "Opening"), col("closing_meter", "Closing"), col("sales_liters", "Sales Litres"), col("rate", "Rate (৳)"), col("total", "Amount (৳)"), col("operator_name", "Operator"), col("notes", "Notes")]
        raw = fetchall("SELECT m.*,sh.name shift_name,p.pump_number,n.nozzle_number,pr.name product_name,e.name operator_name FROM meter_readings m LEFT JOIN shifts sh ON sh.id=m.shift_id JOIN pumps p ON p.id=m.pump_id JOIN nozzles n ON n.id=m.nozzle_id JOIN products pr ON pr.id=m.product_id LEFT JOIN employees e ON e.id=m.operator_id WHERE m.date BETWEEN ? AND ? ORDER BY m.date DESC,m.id DESC", (start, end))
        rows = [{**dict(r), "rate": to_money(r["rate_paisa"]), "total": to_money(r["total_paisa"])} for r in raw]
    elif report_type == "fuel_stock":
        columns = [col("product", "Product"), col("unit", "Unit"), col("opening_stock", "Opening"), col("purchased", "Purchased"), col("sold", "Sold"), col("adjusted", "Adjusted"), col("closing_book_stock", "Closing book"), col("physical_stock", "Last physical"), col("difference", "Difference"), col("physical_verified", "Verified")]
        rows = [r for r in stock_position(start, end) if r["category"] == "fuel"]
    elif report_type == "tank_dips":
        columns = [col("date", "Date"), col("tank_number", "Tank"), col("product_name", "Product"), col("dip_reading", "Dip reading"), col("estimated_quantity", "Estimated quantity"), col("book_stock", "Book stock"), col("physical_stock", "Physical stock"), col("difference", "Difference"), col("verified", "Verified"), col("operator_name", "Operator"), col("notes", "Notes")]
        raw = fetchall("SELECT d.*,t.tank_number,p.name product_name,e.name operator_name FROM tank_dips d JOIN tanks t ON t.id=d.tank_id JOIN products p ON p.id=t.product_id LEFT JOIN employees e ON e.id=d.operator_id WHERE d.date BETWEEN ? AND ? ORDER BY d.date DESC,d.id DESC", (start, end))
        rows = [{**dict(r), "verified": "Yes" if r["verified"] else "Unverified"} for r in raw]
    elif report_type == "purchases":
        columns = [col("invoice_no", "Invoice"), col("date", "Date"), col("supplier_name", "Supplier"), col("product_name", "Product"), col("quantity", "Quantity"), col("purchase_rate", "Rate (৳)"), col("total", "Total (৳)"), col("paid", "Paid (৳)"), col("due", "Due (৳)"), col("challan_no", "Challan")]
        raw = fetchall("SELECT pu.*,s.name supplier_name,p.name product_name FROM purchases pu JOIN suppliers s ON s.id=pu.supplier_id JOIN products p ON p.id=pu.product_id WHERE pu.date BETWEEN ? AND ? ORDER BY pu.date DESC,pu.id DESC", (start, end))
        rows = [{**dict(r), "purchase_rate": to_money(r["purchase_rate_paisa"]), "total": to_money(r["total_paisa"]), "paid": to_money(r["paid_amount_paisa"]), "due": to_money(r["due_amount_paisa"])} for r in raw]
    elif report_type in {"customer_due", "supplier_due"}:
        if report_type == "customer_due":
            columns = [col("name", "Customer"), col("phone", "Phone"), col("vehicle_no", "Vehicle"), col("customer_type", "Type"), col("total_credit", "Credit sales (৳)"), col("total_payments", "Payments (৳)"), col("current_due", "Current due (৳)"), col("last_payment", "Last payment")]
            raw = fetchall("""SELECT c.id,c.name,c.phone,c.vehicle_no,c.customer_type,c.opening_due_paisa,
                COALESCE(SUM(CASE WHEN t.amount_paisa>0 THEN t.amount_paisa ELSE 0 END),0) total_credit,
                COALESCE(SUM(CASE WHEN t.type='payment' THEN -t.amount_paisa ELSE 0 END),0) total_payments,
                COALESCE(MAX(CASE WHEN t.type='payment' THEN t.date END),'') last_payment,
                c.opening_due_paisa+COALESCE(SUM(t.amount_paisa),0) current_due_paisa
                FROM customers c LEFT JOIN customer_transactions t ON t.customer_id=c.id GROUP BY c.id ORDER BY current_due_paisa DESC""")
            rows = [{**dict(r), "total_credit": to_money(r["total_credit"]), "total_payments": to_money(r["total_payments"]), "current_due": to_money(r["current_due_paisa"])} for r in raw]
        else:
            columns = [col("name", "Supplier"), col("company", "Company"), col("phone", "Phone"), col("total_purchases", "Purchases (৳)"), col("total_paid", "Payments (৳)"), col("current_due", "Current due (৳)")]
            raw = fetchall("""SELECT s.id,s.name,s.company,s.phone,s.opening_balance_paisa,
                COALESCE(SUM(CASE WHEN t.amount_paisa>0 THEN t.amount_paisa ELSE 0 END),0) total_purchases,
                COALESCE(SUM(CASE WHEN t.type='payment' THEN -t.amount_paisa ELSE 0 END),0) total_paid,
                s.opening_balance_paisa+COALESCE(SUM(t.amount_paisa),0) current_due_paisa
                FROM suppliers s LEFT JOIN supplier_transactions t ON t.supplier_id=s.id GROUP BY s.id ORDER BY current_due_paisa DESC""")
            rows = [{**dict(r), "total_purchases": to_money(r["total_purchases"]), "total_paid": to_money(r["total_paid"]), "current_due": to_money(r["current_due_paisa"])} for r in raw]
    elif report_type == "customer_statement":
        customer_id = parse_int(args.get("customer_id"), "customer", None)
        if not customer_id:
            raise ApiError("Select a customer for a statement.")
        stmt = fetchone("SELECT * FROM customers WHERE id=?", (customer_id,))
        if not stmt: raise ApiError("Customer not found.", 404)
        columns = [col("date", "Date"), col("type", "Type"), col("description", "Description"), col("debit", "Debit (৳)"), col("credit", "Credit (৳)"), col("balance", "Balance (৳)")]
        trans = fetchall("SELECT * FROM customer_transactions WHERE customer_id=? AND date BETWEEN ? AND ? ORDER BY date,id", (customer_id, start, end))
        before = fetchone("SELECT COALESCE(SUM(amount_paisa),0) val FROM customer_transactions WHERE customer_id=? AND date<?", (customer_id, start))["val"]
        balance = int(stmt["opening_due_paisa"]) + int(before or 0)
        rows.append({"date": start, "type": "Opening balance", "description": stmt["name"], "debit": 0, "credit": to_money(balance), "balance": to_money(balance)})
        for r in trans:
            amount = int(r["amount_paisa"]); balance += amount
            rows.append({"date": r["date"], "type": r["type"], "description": r["notes"] or r["source"] or r["type"], "debit": to_money(max(-amount, 0)), "credit": to_money(max(amount, 0)), "balance": to_money(balance)})
        summary["customer"] = stmt["name"]; summary["current_due"] = to_money(balance)
    elif report_type == "cashbook":
        columns = [col("date", "Date"), col("type", "Type"), col("direction", "Direction"), col("amount", "Amount (৳)"), col("notes", "Description"), col("shift_name", "Shift"), col("user_name", "Entered by")]
        raw = fetchall("SELECT c.*,u.full_name user_name,sh.name shift_name FROM cash_transactions c LEFT JOIN users u ON u.id=c.user_id LEFT JOIN shifts sh ON sh.id=c.shift_id WHERE c.date BETWEEN ? AND ? ORDER BY c.date,c.id", (start, end))
        rows = [{**dict(r), "amount": to_money(r["amount_paisa"])} for r in raw]
        summary["cash_balance"] = to_money(cash_balance())
    elif report_type == "bank_deposits":
        columns = [col("date", "Date"), col("bank_name", "Bank"), col("account_name", "Account name"), col("direction", "Direction"), col("type", "Type"), col("amount", "Amount (৳)"), col("slip_no", "Slip No"), col("depositor", "Depositor"), col("reference", "Reference"), col("shift_name", "Shift")]
        raw = fetchall("SELECT t.*,b.bank_name,b.account_name,sh.name shift_name FROM bank_transactions t LEFT JOIN bank_accounts b ON b.id=t.bank_id LEFT JOIN shifts sh ON sh.id=t.shift_id WHERE t.date BETWEEN ? AND ? AND t.type='cash_deposit' AND t.direction='in' ORDER BY t.date DESC,t.id DESC", (start, end))
        rows = [{**dict(r), "amount": to_money(r["amount_paisa"])} for r in raw]
    elif report_type == "expenses":
        columns = [col("date", "Date"), col("category", "Category"), col("description", "Description"), col("amount", "Amount (৳)"), col("payment_method", "Payment"), col("shift_name", "Shift"), col("paid_to", "Paid to"), col("voucher_no", "Voucher"), col("status", "Status"), col("created_by_name", "Created by")]
        raw = fetchall("SELECT x.*,u.full_name created_by_name,sh.name shift_name FROM expenses x LEFT JOIN users u ON u.id=x.created_by LEFT JOIN shifts sh ON sh.id=x.shift_id WHERE x.date BETWEEN ? AND ? ORDER BY x.date DESC,x.id DESC", (start, end))
        rows = [{**dict(r), "amount": to_money(r["amount_paisa"])} for r in raw]
        summary["approved_total"] = to_money(sum(r["amount_paisa"] for r in raw if r["status"] == "approved"))
    elif report_type == "profit_loss":
        columns = [col("line", "Profit & Loss"), col("amount", "Amount (৳)")]
        revenue, cogs = sales_revenue_and_cogs(start, end)
        expenses = int(fetchone("SELECT COALESCE(SUM(amount_paisa),0) value FROM expenses WHERE date BETWEEN ? AND ? AND status='approved'", (start, end))["value"] or 0)
        operating_income = fetchone("SELECT COALESCE(SUM(amount_paisa),0) value FROM cash_transactions WHERE date BETWEEN ? AND ? AND direction='in' AND type='other_income'", (start, end))["value"]
        gross = revenue - cogs
        net = gross - expenses
        rows = [{"line": "Sales revenue (cash + credit)", "amount": to_money(revenue)}, {"line": "Cost of goods sold · moving weighted average", "amount": to_money(-cogs)}, {"line": "Gross profit", "amount": to_money(gross)}, {"line": "Other operating income", "amount": to_money(operating_income)}, {"line": "Operating expenses", "amount": to_money(-expenses)}, {"line": "Net profit", "amount": to_money(net + int(operating_income or 0))}]
        summary = {"sales": to_money(revenue), "cogs": to_money(cogs), "gross_profit": to_money(gross), "expenses": to_money(expenses), "net_profit": to_money(net + int(operating_income or 0))}
    elif report_type == "employee_activity":
        columns = [col("timestamp", "Date & time"), col("user_name", "User"), col("action", "Action"), col("module", "Module"), col("record_id", "Record ID"), col("reason", "Reason")]
        raw = fetchall("SELECT a.*,u.full_name user_name FROM audit_logs a LEFT JOIN users u ON u.id=a.user_id WHERE substr(a.timestamp,1,10) BETWEEN ? AND ? ORDER BY a.id DESC", (start, end))
        rows = [dict(r) for r in raw]
    elif report_type == "audit_logs":
        columns = [col("timestamp", "Date & time"), col("user_name", "User"), col("action", "Action"), col("module", "Module"), col("record_id", "Record ID"), col("previous_value", "Previous value"), col("new_value", "New value"), col("reason", "Reason")]
        raw = fetchall("SELECT a.*,u.full_name user_name FROM audit_logs a LEFT JOIN users u ON u.id=a.user_id WHERE substr(a.timestamp,1,10) BETWEEN ? AND ? ORDER BY a.id DESC", (start, end))
        rows = []
        for r in raw:
            d = dict(r)
            for key in ("previous_value", "new_value"):
                if d[key]:
                    try: d[key] = json.loads(d[key])
                    except ValueError: pass
            rows.append(d)
    elif report_type == "price_changes":
        columns = [col("effective_at", "Date & time"), col("product_name", "Product"), col("old_price", "Old price (৳)"), col("new_price", "New price (৳)"), col("user_name", "Changed by"), col("reason", "Reason")]
        raw = fetchall("SELECT h.*,p.name product_name,u.full_name user_name FROM price_history h JOIN products p ON p.id=h.product_id LEFT JOIN users u ON u.id=h.user_id WHERE substr(h.effective_at,1,10) BETWEEN ? AND ? ORDER BY h.id DESC", (start, end))
        rows = [{**dict(r), "old_price": to_money(r["old_price_paisa"]), "new_price": to_money(r["new_price_paisa"])} for r in raw]
    else:
        raise ApiError("Select a supported report.", 404, "unknown_report")
    return {"type": report_type, "title": REPORT_TITLES[report_type], "start": start, "end": end, "columns": columns, "rows": rows, "summary": summary}


@app.get("/api/reports")
@permission_required("reports:view")
def reports_api():
    report_type = str(request.args.get("type", "daily_sales"))
    start = parse_date(request.args.get("start"), "start date") if request.args.get("start") else today_text()
    end = parse_date(request.args.get("end"), "end date") if request.args.get("end") else today_text()
    if start > end:
        raise ApiError("Start date must be on or before end date.")
    report = make_report(report_type, start, end, request.args.to_dict())
    report["business"] = {key: settings_dict().get(key, "") for key in ("business_name", "dealer_name", "business_address")}
    report["generated_at"] = now_text()
    report["generated_by"] = g.user["full_name"]
    return jsonify(report)


@app.get("/api/export/<report_type>")
@permission_required("reports:view")
def export_report(report_type: str):
    start = parse_date(request.args.get("start"), "start date") if request.args.get("start") else today_text()
    end = parse_date(request.args.get("end"), "end date") if request.args.get("end") else today_text()
    report = make_report(report_type, start, end, request.args.to_dict())
    export_format = str(request.args.get("format", "csv")).lower()
    filename_base = report_type.replace("_", "-") + f"-{start}-to-{end}"
    if export_format == "csv":
        stream = io.StringIO()
        writer = csv.writer(stream)
        writer.writerow([c["label"] for c in report["columns"]])
        for row in report["rows"]:
            writer.writerow([row.get(c["key"], "") for c in report["columns"]])
        content = ("\ufeff" + stream.getvalue()).encode("utf-8")
        return send_file(io.BytesIO(content), as_attachment=True, download_name=filename_base + ".csv", mimetype="text/csv; charset=utf-8")
    if export_format == "xlsx":
        book = Workbook()
        sheet = book.active
        sheet.title = report["title"][:31]
        sheet.append([setting("business_name"), setting("dealer_name")])
        sheet.append([report["title"]])
        sheet.append([f"{start} — {end}"])
        sheet.append([f"Generated: {now_text()} · {g.user['full_name']}"])
        sheet.append([column["label"] for column in report["columns"]])
        for row in report["rows"]:
            sheet.append([row.get(column["key"], "") if not isinstance(row.get(column["key"], ""), (dict, list)) else json.dumps(row.get(column["key"]), ensure_ascii=False) for column in report["columns"]])
        for cell in sheet[5]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = PatternFill("solid", fgColor="173D72")
            cell.alignment = Alignment(horizontal="center")
        for column_cells in sheet.columns:
            width = min(40, max(12, max((len(str(cell.value or "")) for cell in column_cells), default=10) + 2))
            sheet.column_dimensions[column_cells[0].column_letter].width = width
        sheet.freeze_panes = "A6"
        output = io.BytesIO(); book.save(output); output.seek(0)
        return send_file(output, as_attachment=True, download_name=filename_base + ".xlsx", mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    raise ApiError("Export format must be CSV or XLSX.")

# ---------- Notifications, global search and security backups ----------
@app.get("/api/notifications")
@permission_required("notifications:view")
def notifications_api():
    items = []
    products = fetchall("SELECT id,name,category,unit,current_stock,min_stock FROM products WHERE active=1 ORDER BY name")
    for product in products:
        stock = float(product["current_stock"])
        minimum = float(product["min_stock"])
        if (product["category"] == "fuel" and stock <= minimum) or (minimum > 0 and stock <= minimum):
            items.append({"id": f"stock-{product['id']}", "kind": "low_stock", "priority": "high", "title": f"Low stock · {product['name']}",
                          "message": f"Current stock {stock:,.3f} {product['unit']} (minimum {minimum:,.3f}).", "module": "stock", "date": today_text()})
    threshold = money_minor(setting("due_warning", "50000"))
    for row in fetchall("""SELECT c.id,c.name,c.opening_due_paisa+COALESCE(SUM(t.amount_paisa),0) due FROM customers c
        LEFT JOIN customer_transactions t ON t.customer_id=c.id WHERE c.status='active' GROUP BY c.id
        HAVING due>0 ORDER BY due DESC LIMIT 100"""):
        due = int(row["due"] or 0)
        if due > threshold:
            items.append({"id": f"customer-due-{row['id']}", "kind": "customer_due", "priority": "high" if due >= threshold * 2 else "medium",
                          "title": f"Outstanding customer balance · {row['name']}", "message": f"Current due ৳{to_money(due):,.2f}.", "module": "customer_due", "date": today_text()})
    for row in fetchall("""SELECT s.id,s.name,s.opening_balance_paisa+COALESCE(SUM(t.amount_paisa),0) due FROM suppliers s
        LEFT JOIN supplier_transactions t ON t.supplier_id=s.id WHERE s.status='active' GROUP BY s.id HAVING due>0 ORDER BY due DESC LIMIT 100"""):
        items.append({"id": f"supplier-due-{row['id']}", "kind": "supplier_due", "priority": "medium", "title": f"Supplier due · {row['name']}",
                      "message": f"Outstanding payable ৳{to_money(row['due']):,.2f}.", "module": "suppliers", "date": today_text()})
    for dip in fetchall("""SELECT d.id,d.date,d.difference,d.verified,t.tank_number,p.name product_name,t.capacity FROM tank_dips d
        JOIN tanks t ON t.id=d.tank_id JOIN products p ON p.id=t.product_id WHERE d.verified=0 AND ABS(d.difference)>MAX(10,t.capacity*0.02)
        ORDER BY d.date DESC,d.id DESC LIMIT 50"""):
        items.append({"id": f"dip-{dip['id']}", "kind": "stock_shortage", "priority": "high", "title": f"Tank dip discrepancy · {dip['tank_number']}",
                      "message": f"{dip['product_name']} book variance {float(dip['difference']):,.3f}. Physical reading is unverified.", "module": "tanks", "date": dip["date"]})
    for shift in fetchall("SELECT id,name FROM shifts WHERE active=1"):
        if not fetchone("SELECT 1 FROM shift_closings WHERE date=? AND shift_id=?", (today_text(), shift["id"])):
            items.append({"id": f"shift-close-{shift['id']}", "kind": "missing_shift_close", "priority": "medium", "title": f"Shift closing pending · {shift['name']}",
                          "message": "No shift close has been recorded for today.", "module": "shifts", "date": today_text()})
    for expense in fetchall("SELECT id,category,amount_paisa,date FROM expenses WHERE status='pending' ORDER BY id DESC LIMIT 50"):
        items.append({"id": f"expense-{expense['id']}", "kind": "unapproved_expense", "priority": "medium", "title": f"Expense awaiting approval · {expense['category']}",
                      "message": f"৳{to_money(expense['amount_paisa']):,.2f} · submitted {expense['date']}.", "module": "expenses", "date": expense["date"]})
    cutoff = (now_dt() - timedelta(hours=24)).strftime("%Y-%m-%d %H:%M:%S")
    for change in fetchall("SELECT h.id,p.name,h.old_price_paisa,h.new_price_paisa,h.effective_at FROM price_history h JOIN products p ON p.id=h.product_id WHERE h.effective_at>=? ORDER BY h.id DESC LIMIT 30", (cutoff,)):
        items.append({"id": f"price-{change['id']}", "kind": "price_change", "priority": "low", "title": f"Price changed · {change['name']}",
                      "message": f"৳{to_money(change['old_price_paisa']):,.2f} → ৳{to_money(change['new_price_paisa']):,.2f}.", "module": "prices", "date": change["effective_at"]})
    failed = fetchall("SELECT id,created_at,message FROM backup_history WHERE status='failed' ORDER BY id DESC LIMIT 5")
    for item in failed:
        items.append({"id": f"backup-{item['id']}", "kind": "failed_backup", "priority": "high", "title": "Backup failed", "message": item["message"] or "Review backup configuration.", "module": "backup", "date": item["created_at"]})
    items.sort(key=lambda x: (0 if x["priority"] == "high" else 1 if x["priority"] == "medium" else 2, x["date"]), reverse=False)
    return jsonify({"items": items, "count": len(items)})


@app.get("/api/search")
@auth_required
def global_search():
    term = str(request.args.get("q", "")).strip()
    if len(term) < 2:
        return jsonify({"results": []})
    like = f"%{term}%"
    perms = set(user_permissions(g.user))
    results = []
    if "sales:view" in perms:
        for row in fetchall("SELECT id,memo_no,date,vehicle_no,total_paisa,status FROM sales WHERE memo_no LIKE ? OR vehicle_no LIKE ? ORDER BY id DESC LIMIT 8", (like, like)):
            results.append({"type": "memo", "id": row["id"], "title": f"Memo {row['memo_no']}", "subtitle": f"{row['date']} · {row['vehicle_no'] or 'No vehicle'} · ৳{to_money(row['total_paisa']):,.2f}", "module": "sales_history"})
        for row in fetchall("SELECT DISTINCT s.id,s.memo_no,s.date,s.total_paisa FROM sales s JOIN sale_items si ON si.sale_id=s.id JOIN products p ON p.id=si.product_id WHERE p.name LIKE ? ORDER BY s.id DESC LIMIT 5", (like,)):
            results.append({"type": "product_sale", "id": row["id"], "title": f"{row['memo_no']} · product match", "subtitle": f"{row['date']} · ৳{to_money(row['total_paisa']):,.2f}", "module": "sales_history"})
    if "customers:manage" in perms or "due:collect" in perms:
        for row in fetchall("SELECT id,name,phone,vehicle_no FROM customers WHERE name LIKE ? OR phone LIKE ? OR vehicle_no LIKE ? ORDER BY name LIMIT 8", (like, like, like)):
            results.append({"type": "customer", "id": row["id"], "title": row["name"], "subtitle": " · ".join([x for x in [row["phone"], row["vehicle_no"]] if x]), "module": "customers"})
    if "products:view" in perms:
        for row in fetchall("SELECT id,name,category,current_stock,unit FROM products WHERE name LIKE ? ORDER BY name LIMIT 8", (like,)):
            results.append({"type": "product", "id": row["id"], "title": row["name"], "subtitle": f"{float(row['current_stock']):,.3f} {row['unit']} · {row['category']}", "module": "products"})
    if "purchases:view" in perms:
        for row in fetchall("SELECT pu.id,pu.invoice_no,pu.date,s.name supplier_name,pu.total_paisa FROM purchases pu JOIN suppliers s ON s.id=pu.supplier_id WHERE pu.invoice_no LIKE ? OR pu.challan_no LIKE ? OR s.name LIKE ? ORDER BY pu.id DESC LIMIT 8", (like, like, like)):
            results.append({"type": "purchase", "id": row["id"], "title": row["invoice_no"], "subtitle": f"{row['supplier_name']} · {row['date']} · ৳{to_money(row['total_paisa']):,.2f}", "module": "purchases"})
    if "employees:manage" in perms:
        for row in fetchall("SELECT id,name,designation FROM employees WHERE name LIKE ? OR phone LIKE ? ORDER BY name LIMIT 5", (like, like)):
            results.append({"type": "employee", "id": row["id"], "title": row["name"], "subtitle": row["designation"] or "Employee", "module": "employees"})
    return jsonify({"results": results[:25]})


def create_backup_file(backup_type: str = "manual", user_id: int | None = None) -> dict:
    DATA_DIR.mkdir(parents=True, exist_ok=True); BACKUP_DIR.mkdir(parents=True, exist_ok=True)
    stamp = now_dt().strftime("%Y%m%d-%H%M%S")
    filename = f"hok-{backup_type}-{stamp}-{uuid.uuid4().hex[:4]}.sqlite3"
    path = BACKUP_DIR / filename
    source = get_db()
    dest = sqlite3.connect(str(path))
    try:
        source.backup(dest)
    finally:
        dest.close()
    size = path.stat().st_size
    with atomic() as conn:
        cur = conn.execute("INSERT INTO backup_history(filename,path,status,backup_type,size_bytes,created_by,created_at,message) VALUES(?,?,'success',?,?,?,?,?)",
                           (filename, str(path), backup_type, size, user_id, now_text(), "Database backup completed"))
        if backup_type == "automatic":
            set_setting(conn, "last_auto_backup", now_text())
        audit(conn, "backup_created", "backup", cur.lastrowid, None, {"filename": filename, "size_bytes": size, "type": backup_type})
    return {"id": cur.lastrowid, "filename": filename, "size_bytes": size, "created_at": now_text(), "status": "success"}


@app.get("/api/backups")
@permission_required("backup:manage")
def list_backups():
    rows = fetchall("SELECT b.*,u.full_name AS created_by_name FROM backup_history b LEFT JOIN users u ON u.id=b.created_by ORDER BY b.id DESC LIMIT 500")
    return jsonify([dict(r) | {"download_url": f"/api/backups/{r['id']}/download" if r["status"] == "success" else None} for r in rows])


@app.post("/api/backups")
@permission_required("backup:manage")
def create_backup_api():
    result = create_backup_file("manual", g.user["id"])
    return jsonify(result), 201


@app.get("/api/backups/<int:backup_id>/download")
@permission_required("backup:manage")
def download_backup(backup_id: int):
    row = fetchone("SELECT * FROM backup_history WHERE id=? AND status='success'", (backup_id,))
    if not row:
        raise ApiError("Backup file not found.", 404)
    path = Path(row["path"]).resolve()
    if not path.is_relative_to(BACKUP_DIR.resolve()) or not path.exists():
        raise ApiError("Backup file is unavailable.", 404)
    return send_file(path, as_attachment=True, download_name=row["filename"], mimetype="application/vnd.sqlite3")


@app.post("/api/backups/<int:backup_id>/restore")
@permission_required("backup:manage")
def restore_backup(backup_id: int):
    data = request_data()
    if str(data.get("confirm", "")).strip().upper() != "RESTORE":
        raise ApiError("Type RESTORE to confirm replacing the current database.")
    row = fetchone("SELECT * FROM backup_history WHERE id=? AND status='success'", (backup_id,))
    if not row:
        raise ApiError("Backup file not found.", 404)
    path = Path(row["path"]).resolve()
    if not path.is_relative_to(BACKUP_DIR.resolve()) or not path.exists():
        raise ApiError("Backup file is unavailable.", 404)
    check = sqlite3.connect(str(path))
    try:
        integrity = check.execute("PRAGMA integrity_check").fetchone()[0]
        tables = {r[0] for r in check.execute("SELECT name FROM sqlite_master WHERE type='table'")}
        if integrity != "ok" or not {"users", "sales", "products", "settings"}.issubset(tables):
            raise ApiError("The selected file is not a valid HOK station backup.")
    finally:
        check.close()
    safety = create_backup_file("pre-restore", g.user["id"])
    source = sqlite3.connect(str(path))
    target = get_db()
    try:
        target.execute("PRAGMA wal_checkpoint(TRUNCATE)")
        target.execute("PRAGMA foreign_keys=OFF")
        source.backup(target)
        target.commit()
        target.execute("PRAGMA foreign_keys=ON")
    except sqlite3.Error as exc:
        raise ApiError(f"Restore failed; safety backup {safety['filename']} was kept. {exc}", 500, "restore_failed")
    finally:
        source.close()
    with atomic() as conn:
        audit(conn, "backup_restored", "backup", backup_id, {"current_database": "before restore"}, {"filename": row["filename"], "safety_backup": safety["filename"]})
    return jsonify({"ok": True, "safety_backup": safety["filename"], "message": "Database restored successfully. Please sign in again if your user changed."})


@app.get("/print/memo/<int:sale_id>")
@permission_required("sales:view")
def print_memo(sale_id: int):
    sale = sale_json(sale_id)
    products = fetchall("SELECT id,name,category,unit FROM products WHERE active=1 ORDER BY id")
    return render_template("memo_print.html", sale=sale, products=products, settings=settings_dict())


@app.get("/print/report")
@permission_required("reports:view")
def print_report():
    report_type = str(request.args.get("type", "daily_sales"))
    start = parse_date(request.args.get("start"), "start date") if request.args.get("start") else today_text()
    end = parse_date(request.args.get("end"), "end date") if request.args.get("end") else today_text()
    report = make_report(report_type, start, end, request.args.to_dict())
    return render_template("report_print.html", report=report, settings=settings_dict(), user=g.user, generated_at=now_text())


@app.get("/api/stock")
@permission_required("inventory:view")
def get_stock():
    start = parse_date(request.args.get("start"), "start date") if request.args.get("start") else today_text()
    end = parse_date(request.args.get("end"), "end date") if request.args.get("end") else today_text()
    return jsonify(stock_position(start, end))


@app.get("/api/tank-dips")
@permission_required("inventory:view")
def list_tank_dips():
    start = parse_date(request.args.get("start"), "start date") if request.args.get("start") else "0001-01-01"
    end = parse_date(request.args.get("end"), "end date") if request.args.get("end") else today_text()
    rows = fetchall("""SELECT d.*,t.tank_number,p.name product_name,e.name operator_name,u.full_name verified_by_name
        FROM tank_dips d JOIN tanks t ON t.id=d.tank_id JOIN products p ON p.id=t.product_id
        LEFT JOIN employees e ON e.id=d.operator_id LEFT JOIN users u ON u.id=d.verified_by
        WHERE d.date BETWEEN ? AND ? ORDER BY d.date DESC,d.id DESC LIMIT 2000""", (start, end))
    return jsonify([dict(r) for r in rows])


@app.get("/api/audit-logs")
@permission_required("audit:view")
def list_audit_logs():
    start = parse_date(request.args.get("start"), "start date") if request.args.get("start") else "0001-01-01"
    end = parse_date(request.args.get("end"), "end date") if request.args.get("end") else today_text()
    sql = "SELECT a.*,u.full_name AS user_name FROM audit_logs a LEFT JOIN users u ON u.id=a.user_id WHERE substr(a.timestamp,1,10) BETWEEN ? AND ?"
    params: list[Any] = [start, end]
    if request.args.get("module"):
        sql += " AND a.module=?"; params.append(request.args["module"])
    term = str(request.args.get("q", "")).strip()
    if term:
        sql += " AND (a.action LIKE ? OR a.module LIKE ? OR a.record_id LIKE ? OR u.full_name LIKE ? OR a.reason LIKE ?)"
        params.extend([f"%{term}%"] * 5)
    sql += " ORDER BY a.id DESC LIMIT 2000"
    return jsonify([audit_dict(row) for row in fetchall(sql, tuple(params))])


@app.get("/api/suppliers/<int:supplier_id>/statement")
@permission_required("reports:view")
def supplier_statement(supplier_id: int):
    supplier = active_record("suppliers", supplier_id)
    start = parse_date(request.args.get("start"), "start date") if request.args.get("start") else "0001-01-01"
    end = parse_date(request.args.get("end"), "end date") if request.args.get("end") else today_text()
    opening_tx = fetchone("SELECT COALESCE(SUM(amount_paisa),0) value FROM supplier_transactions WHERE supplier_id=? AND date<?", (supplier_id, start))["value"]
    balance = int(supplier["opening_balance_paisa"]) + int(opening_tx or 0)
    rows = [{"date": start if start != "0001-01-01" else "", "type": "Opening balance", "description": supplier["name"], "debit": 0, "credit": to_money(balance), "balance": to_money(balance)}]
    transactions = fetchall("SELECT * FROM supplier_transactions WHERE supplier_id=? AND date BETWEEN ? AND ? ORDER BY date,id", (supplier_id, start, end))
    for entry in transactions:
        amount = int(entry["amount_paisa"]); balance += amount
        rows.append({"date": entry["date"], "type": entry["type"], "description": entry["notes"] or entry["source"] or entry["type"], "debit": to_money(max(-amount, 0)), "credit": to_money(max(amount, 0)), "balance": to_money(balance)})
    return jsonify({"supplier": dict(supplier), "rows": rows, "current_due": to_money(balance)})


@app.put("/api/products/<int:product_id>/price")
@permission_required("prices:manage")
def change_product_price(product_id: int):
    product = active_record("products", product_id)
    data = request_data()
    new_price = money_minor(data.get("selling_price"), "selling price")
    reason = str(data.get("reason", "")).strip()
    if not reason:
        raise ApiError("A reason is required for every selling-price change.")
    if new_price == int(product["selling_price_paisa"]):
        return jsonify({"ok": True, "changed": False})
    with atomic() as conn:
        conn.execute("UPDATE products SET selling_price_paisa=?,updated_at=? WHERE id=?", (new_price, now_text(), product_id))
        conn.execute("INSERT INTO price_history(product_id,old_price_paisa,new_price_paisa,effective_at,user_id,reason) VALUES(?,?,?,?,?,?)",
                     (product_id, product["selling_price_paisa"], new_price, now_text(), g.user["id"], reason))
        audit(conn, "price_changed", "prices", product_id, {"product": product["name"], "selling_price": to_money(product["selling_price_paisa"])}, {"product": product["name"], "selling_price": to_money(new_price)}, reason)
    return jsonify({"ok": True, "changed": True, "selling_price": to_money(new_price)})


@app.errorhandler(sqlite3.IntegrityError)
def handle_integrity_error(error):
    app.logger.warning("Database constraint: %s", error)
    if request.path.startswith("/api/"):
        return jsonify({"error": "This record conflicts with existing data. Check for a duplicate number or linked record.", "code": "constraint_violation"}), 409
    return "A data constraint prevented this request.", 409


@app.errorhandler(Exception)
def handle_unexpected_error(error):
    if isinstance(error, ApiError):
        return handle_api_error(error)
    app.logger.exception("Unhandled application error")
    if request.path.startswith("/api/"):
        return jsonify({"error": "An unexpected server error occurred. No changes were committed.", "code": "internal_error"}), 500
    return "An unexpected error occurred.", 500


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=int(os.environ.get("PORT", 5000)), debug=os.environ.get("HOK_DEBUG", "0") == "1")
